"""python3 -m unittest discover -s excel_filter"""

import io
import json
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

import excel_filter
from excel_filter import (detect_columns, distinct, download, load, production_summary, read,
                          sentences_with, summarize)

HEADER = ("Data", "Dia", "Horário", "Hora", "Turno", "Plan", "Produzido", "Principais perdas", "Vazia")
ROWS = [
    (date(2026, 1, 10), 10, "06h - 07h", 1, 1, 1885, 1700, "Monobloco - parada de linha pela qualidade", None),
    (date(2026, 2, 20), 20, "07h - 08h", 2, 2, 1885, 1900, "Tampadora: travamento de tampa no mandril", None),
    ("05/03/2026", 5, "06h - 07h", 3, 3, 1885, 0, "Gasser - alarme nivel 1, e a tampadora ok", None),
    (date(2026, 4, 1), 1, "08h - 09h", 4, 1, 1885, 1750, None, None),
]


def column(result, name):
    """Values of one column of the summary — the summary drops empty columns."""
    i = result["columns"].index(name)
    return [row[i] for row in result["rows"]]


def workbook(tmp, rows=ROWS):
    book = Workbook()
    sheet = book.active
    sheet.title = "Dados"
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    report = book.create_sheet("Relatório")  # sheet with no header, like the real file
    report["B2"] = "LINHA D02 - 1º TURNO"
    path = Path(tmp) / "orig.xlsx"
    book.save(path)
    return path


class Reading(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.path = workbook(self.tmp)

    def test_picks_the_table_sheet_and_the_three_columns_on_its_own(self):
        sheet, headers, rows = read(self.path)
        self.assertEqual((sheet, len(rows)), ("Dados", 4))
        # "Hora" (1..4) must not win over "Horário" ("06h - 07h")
        self.assertEqual(detect_columns(headers, rows), ("Data", "Horário", "Principais perdas"))

    def test_hour_picker_values(self):
        _, headers, rows = read(self.path)
        self.assertEqual(distinct(headers, rows, "Horário"), ["06h - 07h", "07h - 08h", "08h - 09h"])


class Summary(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        _, self.headers, self.rows = read(workbook(self.tmp))

    def run_summary(self, **kw):
        kw.setdefault("date_col", "Data")
        kw.setdefault("hour_col", "Horário")
        kw.setdefault("text_col", "Principais perdas")
        return summarize(self.headers, self.rows, **kw)

    def test_no_filter_summarises_everything(self):
        result = self.run_summary()
        self.assertEqual((result["found"], result["read"]), (4, 4))
        self.assertEqual(result["period"], (date(2026, 1, 10), date(2026, 4, 1)))

    def test_shift_goes_by_the_hour_the_row_starts_at(self):
        self.assertEqual(self.run_summary(shift=1)["found"], 4)  # 06h, 07h, 06h, 08h
        self.assertEqual(self.run_summary(shift=2)["found"], 0)

    def test_shift_of_each_hour(self):
        said = ("06h - 07h", "13h - 14h", "14h - 15h", "21:00", "22h - 23h", "00h - 01h",
                "05h - 06h", "", None, "sem hora")
        self.assertEqual([excel_filter.shift_of(h) for h in said],
                         [1, 1, 2, 2, 3, 3, 3, None, None, None])

    def test_keyword_is_case_insensitive_and_ors_terms(self):
        self.assertEqual(self.run_summary(terms=["TAMPADORA"])["found"], 2)
        self.assertEqual(self.run_summary(terms=["monobloco", "gasser"])["found"], 2)

    def test_date_range_is_inclusive_and_reads_text_dates(self):
        result = self.run_summary(start=date(2026, 2, 20), end=date(2026, 3, 5))
        self.assertEqual(result["found"], 2)
        self.assertEqual(result["period"], (date(2026, 2, 20), date(2026, 3, 5)))

    def test_three_filters_together(self):
        result = self.run_summary(terms=["tampadora"], start=date(2026, 3, 1),
                                  end=date(2026, 12, 31), shift=1)
        self.assertEqual(column(result, "Principais perdas"),
                         ["Gasser - alarme nivel 1, e a tampadora ok"])

    def test_only_rows_that_have_an_occurrence(self):
        self.assertEqual(self.run_summary(only_with_text=True)["found"], 3)  # a 4ª não tem texto

    def test_totals_add_quantities_and_skip_codes(self):
        # Dia, Hora e Turno são código, não quantidade
        self.assertEqual(self.run_summary()["totals"], {"Plan": 7540, "Produzido": 5350})

    def test_numbers_stored_as_text_are_summed_and_unreadable_cells_are_flagged(self):
        headers = ["Data", "Produzido"]
        rows = [(2, [date(2026, 1, 1), 1000]), (3, [date(2026, 1, 2), "0"]),
                (4, [date(2026, 1, 3), 500]), (5, [date(2026, 1, 4), 700]),
                (6, [date(2026, 1, 5), "sem apontamento"])]
        result = summarize(headers, rows, date_col="Data")
        self.assertEqual(result["totals"]["Produzido"], 2200)  # o "0" texto entra
        self.assertEqual(result["partial"], {"Produzido"})  # o ilegível não vira zero calado

    def test_empty_columns_are_dropped(self):
        # esse dia não tem ocorrência escrita: some "Principais perdas" e "Vazia"
        self.assertEqual(self.run_summary(start=date(2026, 4, 1))["columns"],
                         ["Data", "Dia", "Horário", "Hora", "Turno", "Plan", "Produzido"])

    def test_no_results(self):
        result = self.run_summary(terms=["xyz"])
        self.assertEqual((result["found"], result["rows"], result["totals"], result["period"]),
                         (0, [], {}, None))

    def test_errors(self):
        with self.assertRaises(ValueError):  # inverted range
            self.run_summary(start=date(2026, 5, 1), end=date(2026, 1, 1))
        with self.assertRaises(ValueError):  # keyword with no text column in the sheet
            self.run_summary(text_col="", terms=["a"])
        with self.assertRaises(ValueError):  # column not in the header
            self.run_summary(text_col="Nao existe", terms=["a"])
        broken = ROWS + [("trinta de maio", 30, "09h - 10h", 5, 1, 0, 0, "x", None)]
        _, headers, rows = read(workbook(self.tmp, broken))
        with self.assertRaises(ValueError):  # an unreadable date must not vanish silently
            summarize(headers, rows, date_col="Data", start=date(2026, 1, 1))


class Download(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        _, self.headers, self.rows = read(workbook(self.tmp))
        self.dst = Path(self.tmp) / "ocorrencias.xlsx"

    def written(self, rows, terms=(), line_col=""):
        download(self.dst, self.headers, rows, line_col, "Data", "Horário",
                 "Principais perdas", terms)
        return [[c.value for c in row] for row in load_workbook(self.dst).active.iter_rows()]

    def test_four_columns_one_row_per_occurrence(self):
        found = summarize(self.headers, self.rows, text_col="Principais perdas",
                          only_with_text=True)["matched"]
        sheet = self.written(found)
        self.assertEqual(sheet[0], ["Linha", "Data", "Hora", "Ocorrência"])
        self.assertEqual(len(sheet), 4)  # cabeçalho + as 3 que têm texto
        self.assertEqual(sheet[3][1:3], [datetime(2026, 3, 5), "06h - 07h"])  # data em texto vira data

    def test_keyword_keeps_only_its_sentence(self):
        cell = [(2, [date(2026, 1, 1), 1, "06h - 07h", 1, 1, 0, 0,
                     "Gasser parou. Tampadora travou no mandril.", None])]
        self.assertEqual(self.written(cell, ["tampadora"])[1][3], "Tampadora travou no mandril.")

    def test_a_column_the_sheet_lacks_comes_out_blank(self):
        self.assertIsNone(self.written(self.rows[:1])[1][0])  # esta planilha não tem Linha


if __name__ == "__main__":
    unittest.main()


class ShiftReport(unittest.TestCase):
    """The real workbook keeps one sheet per day; its "Tabela de Dados" has broken dates."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def report(self, days=tuple(range(1, 8)) + (30,), planning=True, header_row=26,
               blocks=(2, 35, 68)):
        book = Workbook()
        book.remove(book.active)
        if planning:
            plan = book.create_sheet("Plan.Prod")
            plan["B1"] = "DATA"
            for i in range(30):
                plan.cell(2 + i, 2).value = datetime(2023, 11, 1 + i)
        broken = book.create_sheet("Tabela de Dados")  # drifting dates, must not win
        broken["A1"] = "Data"
        for i in range(10):
            broken.cell(2 + i, 1).value = datetime(2023, 9, 1 + i % 5)
        for day in days:
            sheet = book.create_sheet(str(day))
            for shift, base in enumerate(blocks, 1):
                for label, offset in (("Horário", 0), ("Formato", 2), ("Plan", 4),
                                      ("Produzido", 6), ("Apontado", 8),
                                      ("Principais Perdas", 12)):
                    sheet.cell(header_row, base + offset).value = label
                for h in range(8):
                    row = header_row + 1 + h
                    sheet.cell(row, base).value = f"{6 * shift + h:02d}h - {6 * shift + h + 1:02d}h"
                    sheet.cell(row, base + 4).value = 1885
                    sheet.cell(row, base + 6).value = 1700
                    if h == 0:
                        sheet.cell(row, base + 12).value = (
                            f"Gasser - alarme no dia {day} turno {shift}.\n"
                            "Tampadora - troca de bico.")
        path = Path(self.tmp) / "turno.xlsx"
        book.save(path)
        return path

    def test_reads_the_day_sheets_and_rebuilds_the_date_from_the_sheet_name(self):
        data = load(self.report())
        self.assertEqual(data["kind"], "turno")
        self.assertEqual(data["month"], (2023, 11))  # planning sheet wins over the broken table
        self.assertEqual(data["month_from"], "Plan.Prod")
        self.assertEqual(len(data["rows"]), 8 * 3 * 8)  # 8 dias x 3 turnos x 8 horas
        self.assertEqual([r[0] for _, r in data["rows"]][:1], [date(2023, 11, 1)])
        self.assertEqual(max(r[0] for _, r in data["rows"]), date(2023, 11, 30))

    def test_a_day_the_month_does_not_have_is_skipped(self):
        dias = tuple(range(1, 8)) + (31,)  # novembro não tem dia 31
        data = load(self.report(days=dias), month=(2023, 11))
        dia = data["headers"].index("Dia")
        self.assertEqual({r[dia] for _, r in data["rows"]}, set(range(1, 8)))

    def test_month_can_be_overridden(self):
        data = load(self.report(), month=(2024, 2))
        self.assertEqual(min(r[0] for _, r in data["rows"]), date(2024, 2, 1))

    def test_another_version_of_the_model_is_read_the_same(self):
        # fevereiro do pacote real: cabeçalho na linha 27, blocos em B/AJ/BR
        outra = load(self.report(header_row=27, blocks=(2, 36, 70)))
        igual = load(self.report())
        self.assertEqual(outra["kind"], "turno")
        self.assertEqual(len(outra["rows"]), len(igual["rows"]))

    def test_without_a_calendar_sheet_it_still_reads_but_says_the_month_is_a_guess(self):
        data = load(self.report(planning=False))
        self.assertIsNone(data["month_from"])  # nenhuma aba merecia nomear o mês
        self.assertTrue(data["month_guessed"])
        self.assertEqual(data["month"], (date.today().year, date.today().month))
        self.assertEqual(len(data["rows"]), 8 * 3 * 8)  # e lê tudo assim mesmo

    def test_a_month_that_was_read_or_asked_for_is_never_marked_as_a_guess(self):
        self.assertFalse(load(self.report())["month_guessed"])
        self.assertFalse(load(self.report(planning=False), month=(2023, 11))["month_guessed"])


class Sentences(unittest.TestCase):
    TEXT = ("Monobloco - linha parada pela qualidade.\n"
            "Gasser - alarme de roda de vácuo.\n"
            "Tampadora: queda de lata na esteira.")

    def test_only_the_sentence_carrying_the_word(self):
        self.assertEqual(sentences_with(self.TEXT, ["gasser"]), "Gasser - alarme de roda de vácuo.")

    def test_two_words_bring_both_sentences(self):
        self.assertEqual(sentences_with(self.TEXT, ["gasser", "tampadora"]).count("\n"), 1)

    def test_without_words_the_whole_text_stays(self):
        self.assertEqual(sentences_with(self.TEXT, []), self.TEXT)

    def test_word_only_in_the_cell_but_in_no_sentence_falls_back_to_everything(self):
        self.assertEqual(sentences_with("uma frase sem ponto", ["nada"]), "uma frase sem ponto")


class History(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        excel_filter.HISTORY_FILE = Path(self.tmp) / "historico.json"

    def test_remembers_newest_first_without_duplicates(self):
        self.assertEqual(excel_filter.history(), [])
        excel_filter.remember("/tmp/a.xlsx", {"rows": 720, "month": [2023, 11]})
        excel_filter.remember("/tmp/b.xlsx", {"rows": 10, "month": None})
        excel_filter.remember("/tmp/a.xlsx", {"rows": 720, "month": [2023, 11]})
        saved = excel_filter.history()
        self.assertEqual([e["path"] for e in saved], ["/tmp/a.xlsx", "/tmp/b.xlsx"])
        self.assertEqual(saved[0]["rows"], 720)

    def test_keeps_only_the_last_entries(self):
        for i in range(excel_filter.HISTORY_LIMIT + 5):
            excel_filter.remember(f"/tmp/{i}.xlsx", {"rows": i, "month": None})
        self.assertEqual(len(excel_filter.history()), excel_filter.HISTORY_LIMIT)

    def test_a_corrupt_file_is_no_history_instead_of_a_crash(self):
        excel_filter.HISTORY_FILE.write_text("{isso não é json", encoding="utf-8")
        self.assertEqual(excel_filter.history(), [])


class MergedOccurrences(unittest.TestCase):
    """In the real sheet one losses cell is merged over several hour rows."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        header_row, base = 26, 2
        book = Workbook()
        book.remove(book.active)
        plan = book.create_sheet("Plan.Prod")
        for i in range(30):
            plan.cell(1 + i, 2).value = datetime(2023, 11, 1 + i)
        for day in range(1, 8):
            sheet = book.create_sheet(str(day))
            sheet.cell(header_row, base).value = "Horário"
            sheet.cell(header_row, base + 12).value = "Principais Perdas"
            for h in range(8):
                sheet.cell(header_row + 1 + h, base).value = f"{6 + h:02d}h - {7 + h:02d}h"
            losses = base + 12
            sheet.cell(header_row + 1, losses).value = "Monobloco - parada longa."
            sheet.merge_cells(start_row=header_row + 1, end_row=header_row + 4,
                              start_column=losses, end_column=losses + 2)
            sheet.cell(header_row + 5, losses).value = "Gasser - alarme."
        path = Path(self.tmp) / "m.xlsx"
        book.save(path)
        self.data = load(path)

    def rows_with_text(self):
        i = self.data["headers"].index("Principais perdas")
        return [r for _, r in self.data["rows"] if str(r[i] or "").strip()]

    def test_the_span_of_a_merged_cell_becomes_the_period(self):
        first = self.rows_with_text()[0]
        headers = self.data["headers"]
        self.assertEqual(first[headers.index("Horário")], "06h - 07h")
        self.assertEqual(first[headers.index("Período")], "06h - 10h")  # 4 horas mescladas
        self.assertEqual(first[headers.index("Horas")].count("|"), 3)

    def test_filtering_by_shift_keeps_a_merged_occurrence_whole(self):
        def found(shift):
            return summarize(self.data["headers"], self.data["rows"], only_with_text=True,
                             text_col="Principais perdas", hour_col="Horas", shift=shift)["found"]
        self.assertEqual(found(1), 14)  # duas por dia, as duas no 1º turno
        self.assertEqual(found(2), 0)

    def test_hours_not_covered_keep_their_own_row_for_the_totals(self):
        self.assertEqual(len(self.data["rows"]), 7 * 8)  # 7 dias x 8 horas do 1º turno


class Lines(unittest.TestCase):
    """A report can hold more than one production line; the user picks which ones."""

    HEADERS = ["Data", "Linha", "Período", "Principais perdas"]
    ROWS = [
        (2, [date(2023, 11, 1), "D02", "06h - 07h", "Gasser - alarme."]),
        (3, [date(2023, 11, 1), "D03", "06h - 07h", "Tampadora - queda de lata."]),
        (4, [date(2023, 11, 2), "D02", "08h - 09h", "Encaixotadora - garra solta."]),
        (5, [date(2023, 11, 2), "D04", "08h - 09h", "Monobloco - parada."]),
    ]

    def run_summary(self, lines=()):
        return summarize(self.HEADERS, self.ROWS, text_col="Principais perdas",
                         date_col="Data", line_col="Linha", lines=lines, only_with_text=True)

    def test_no_line_chosen_means_every_line(self):
        self.assertEqual(self.run_summary()["found"], 4)

    def test_one_line(self):
        result = self.run_summary(["D02"])
        self.assertEqual(column(result, "Linha"), ["D02", "D02"])

    def test_two_lines_at_once(self):
        result = self.run_summary(["D02", "D04"])
        self.assertEqual(sorted(set(column(result, "Linha"))), ["D02", "D04"])
        self.assertEqual(result["found"], 3)

    def test_rows_come_grouped_by_line_so_the_screen_can_block_them(self):
        self.assertEqual(column(self.run_summary(), "Linha"), ["D02", "D02", "D03", "D04"])

    def test_line_names_available_for_the_picker(self):
        self.assertEqual(distinct(self.HEADERS, self.ROWS, "Linha"), ["D02", "D03", "D04"])

    def test_asking_for_a_line_a_sheet_without_lines_cannot_answer(self):
        with self.assertRaises(ValueError):
            summarize(["Data", "Texto"], [(2, [date(2023, 11, 1), "x"])],
                      text_col="Texto", lines=["D02"])

    def test_the_real_report_exposes_its_line(self):
        from excel_filter import LINE_RE
        self.assertEqual(LINE_RE.search("LINHA D02 - 1º TURNO").group(1), "D02")


class Chips(unittest.TestCase):
    """The picked lines show up as chips with an x; needs a display, skipped without one."""

    def setUp(self):
        tk = __import__("tkinter")
        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        self.root.withdraw()
        self.frame = tk.Frame(self.root)
        self.removed = []

    def tearDown(self):
        if hasattr(self, "root"):
            self.root.destroy()

    def test_no_line_picked_says_all_lines(self):
        self.assertEqual(excel_filter.draw_chips(self.frame, [], self.removed.append), [])
        self.assertEqual(self.frame.winfo_children()[0]["text"], "todas as linhas")

    def test_one_chip_per_line_and_the_x_removes_it(self):
        buttons = excel_filter.draw_chips(self.frame, ["D01", "D02"], self.removed.append)
        self.assertEqual(len(buttons), 2)
        buttons[1].invoke()
        self.assertEqual(self.removed, ["D02"])

    def test_redrawing_replaces_the_old_chips(self):
        excel_filter.draw_chips(self.frame, ["D01", "D02"], self.removed.append)
        excel_filter.draw_chips(self.frame, ["D01"], self.removed.append)
        self.assertEqual(len(self.frame.winfo_children()), 1)


class Board(unittest.TestCase):
    """The numbers behind the second tab."""

    HEADERS = ["Data", "SKU", "Plan", "Produzido", "Principais perdas"]
    ROWS = [
        (2, [date(2023, 11, 1), "64846773 - Dove Original", 1885, 1700,
             "Tampadora - alarme de tampa.\nGasser: alarme de roda de vácuo."]),
        (3, [date(2023, 11, 1), "Dove Aer Ap Original - 64846773", 1885, "0",
             "Tampadora – queda de lata."]),
        (4, [date(2023, 11, 2), "64830874 - Rexona Men", 1885, 1000, ""]),
    ]

    def board(self):
        return production_summary(self.HEADERS, self.ROWS)

    def test_planned_produced_and_what_is_missing(self):
        result = self.board()
        self.assertEqual((result["plan"], result["produzido"], result["faltou"]),
                         (5655, 2700, 2955))  # o "0" texto conta como zero produzido
        self.assertAlmostEqual(result["atingido"], 2700 / 5655)

    def test_one_point_per_day_in_date_order(self):
        self.assertEqual(self.board()["dias"],
                         [(date(2023, 11, 1), 3770, 1700), (date(2023, 11, 2), 1885, 1000)])

    def test_each_line_of_the_log_counts_as_one_occurrence_of_its_machine(self):
        self.assertEqual(self.board()["equipamentos"], [("Tampadora", 2), ("Gasser", 1)])

    def test_the_sku_code_groups_the_different_wordings(self):
        codes = [(code, hours) for code, _, hours in self.board()["skus"]]
        self.assertEqual(codes, [("64846773", 2), ("64830874", 1)])  # duas grafias, um código

    def test_an_unreadable_quantity_marks_the_total_as_partial(self):
        rows = self.ROWS + [(5, [date(2023, 11, 3), "", "não apontado", 10, ""])]
        self.assertFalse(production_summary(self.HEADERS, rows)["completo"])

    def test_empty_filter_result_does_not_blow_up(self):
        result = production_summary(self.HEADERS, [])
        self.assertEqual((result["plan"], result["atingido"], result["dias"]), (0, None, []))


class Charts(unittest.TestCase):
    """The canvas drawings; skipped without a display."""

    def setUp(self):
        tk = __import__("tkinter")
        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        self.root.withdraw()
        self.canvas = tk.Canvas(self.root, width=400, height=200)

    def tearDown(self):
        if hasattr(self, "root"):
            self.root.destroy()

    def made(self):
        """The production bars of the combo chart, in drawing order."""
        c = excel_filter.BOARD_COLOURS
        return [i for i in self.canvas.find_withtag("slot") if self.canvas.type(i) == "rectangle"
                and self.canvas.itemcget(i, "fill") in (c["hit"], c["miss"], c["unknown"])]

    def plan_ticks(self):
        return [i for i in self.canvas.find_withtag("slot") if self.canvas.type(i) == "line"]

    def test_combo_chart_draws_one_production_bar_per_day_with_its_plan_tick(self):
        excel_filter.draw_combo_chart(self.canvas, [(date(2023, 11, 1), 10, 1000, 800),
                                                    (date(2023, 11, 2), 0, 1000, 900)])
        self.assertEqual((len(self.made()), len(self.plan_ticks())), (2, 2))

    def test_changeover_is_the_bar_on_top(self):
        excel_filter.draw_combo_chart(self.canvas, [(date(2023, 11, 1), 30, 1000, 800),
                                                    (date(2023, 11, 2), 0, 1000, 900)])
        swap = [i for i in self.canvas.find_withtag("slot") if self.canvas.type(i) == "rectangle"
                and self.canvas.itemcget(i, "fill") == excel_filter.BOARD_COLOURS["swap"]]
        self.assertEqual(len(swap), 1)  # o dia sem troca não ganha barra
        self.assertLess(self.canvas.coords(swap[0])[3], self.canvas.coords(self.made()[0])[1])

    def test_the_bar_is_blue_when_it_reaches_the_plan_and_orange_when_it_does_not(self):
        excel_filter.draw_combo_chart(self.canvas, [(date(2023, 11, 1), 0, 1000, 800),
                                                    (date(2023, 11, 2), 0, 1000, 1200)])
        c = excel_filter.BOARD_COLOURS
        self.assertEqual([self.canvas.itemcget(i, "fill") for i in self.made()],
                         [c["miss"], c["hit"]])

    def test_every_slot_is_tagged_so_the_mouse_finds_its_numbers(self):
        excel_filter.draw_combo_chart(self.canvas, [(date(2023, 11, 1), 5, 1000, 800),
                                                    (date(2023, 11, 2), 5, 1000, 900)])
        self.assertTrue(self.canvas.find_withtag("slot0"))
        self.assertTrue(self.canvas.find_withtag("slot1"))

    def test_the_picked_day_gets_a_highlight_behind_it(self):
        columns = [(date(2023, 11, d), 5, 1000, 900) for d in (1, 2, 3)]
        select = excel_filter.BOARD_COLOURS["select"]

        def highlights():
            return [i for i in self.canvas.find_all() if self.canvas.type(i) == "rectangle"
                    and self.canvas.itemcget(i, "fill") == select]
        excel_filter.draw_combo_chart(self.canvas, columns)
        self.assertEqual(highlights(), [])
        excel_filter.draw_combo_chart(self.canvas, columns, selected=1)
        box = self.canvas.coords(highlights()[0])
        middle = self.canvas.coords(self.made()[1])
        self.assertTrue(box[0] <= middle[0] and middle[2] <= box[2])

    def test_the_tip_shows_the_lines_it_was_given(self):
        excel_filter.draw_tip(self.canvas, 50, 50, ["06h - 07h", "Produzido: 1.700"])
        shown = [self.canvas.itemcget(i, "text") for i in self.canvas.find_withtag("tip")
                 if self.canvas.type(i) == "text"]
        self.assertEqual(shown, ["06h - 07h", "Produzido: 1.700"])

    def test_a_new_tip_replaces_the_old_one(self):
        excel_filter.draw_tip(self.canvas, 50, 50, ["primeiro"])
        excel_filter.draw_tip(self.canvas, 90, 50, ["segundo"])
        shown = [self.canvas.itemcget(i, "text") for i in self.canvas.find_withtag("tip")
                 if self.canvas.type(i) == "text"]
        self.assertEqual(shown, ["segundo"])

    def test_the_tip_stays_inside_the_chart(self):
        # a janela de teste não é mapeada, então vale a mesma largura efetiva que o desenho usa
        efetiva = max(self.canvas.winfo_width(), 200)
        excel_filter.draw_tip(self.canvas, efetiva - 5, 50, ["um texto bem comprido no limite"])
        box = next(i for i in self.canvas.find_withtag("tip")
                   if self.canvas.type(i) == "rectangle")
        self.assertLessEqual(self.canvas.coords(box)[2], efetiva)

    def test_producing_above_the_plan_draws_the_bar_past_the_plan_tick(self):
        excel_filter.draw_combo_chart(self.canvas, [(date(2023, 11, 1), 0, 1000, 1200)])
        bar, line = self.made()[0], self.plan_ticks()[0]
        self.assertLess(self.canvas.coords(bar)[1], self.canvas.coords(line)[1])  # y menor = mais alto

    def painted(self):
        return [i for i in self.canvas.find_all()
                if self.canvas.type(i) == "rectangle" and self.canvas.itemcget(i, "fill")]

    def test_rank_chart_draws_one_bar_per_machine(self):
        excel_filter.draw_rank_chart(self.canvas, [("Tampadora", 270), ("Gasser", 182)])
        self.assertEqual(len(self.painted()), 2)

    def test_every_row_is_tagged_so_the_mouse_can_find_its_machine(self):
        excel_filter.draw_rank_chart(self.canvas, [("Tampadora", 270), ("Gasser", 182)])
        self.assertTrue(self.canvas.find_withtag("bar0"))
        self.assertTrue(self.canvas.find_withtag("bar1"))
        # the clickable strip spans the whole row, not just the painted bar
        strip = max(self.canvas.bbox(i)[2] for i in self.canvas.find_withtag("bar1"))
        bar = max(self.canvas.bbox(i)[2] for i in self.painted())
        self.assertGreater(strip, bar)

    def test_the_highlighted_bar_is_painted_darker(self):
        excel_filter.draw_rank_chart(self.canvas, [("Tampadora", 270), ("Gasser", 182)])
        plain = self.canvas.itemcget(self.painted()[0], "fill")
        excel_filter.draw_rank_chart(self.canvas, [("Tampadora", 270), ("Gasser", 182)], highlight=0)
        self.assertNotEqual(self.canvas.itemcget(self.painted()[0], "fill"), plain)



    def test_nothing_to_draw_says_so_instead_of_an_empty_box(self):
        excel_filter.draw_combo_chart(self.canvas, [])
        texts = [self.canvas.itemcget(i, "text") for i in self.canvas.find_all()]
        self.assertIn("Sem dados no filtro", texts)


class TypingDates(unittest.TestCase):
    """The date fields put the slashes in and accept what people actually type."""

    def test_the_slashes_appear_while_typing(self):
        typed = ["0", "06", "061", "0611", "061120", "06112023"]
        self.assertEqual([excel_filter.mask_date(t) for t in typed],
                         ["0", "06", "06/1", "06/11", "06/11/20", "06/11/2023"])

    def test_it_never_grows_past_a_date(self):
        self.assertEqual(excel_filter.mask_date("0611202345"), "06/11/2023")

    def test_already_formatted_text_is_left_alone(self):
        self.assertEqual(excel_filter.mask_date("06/11/2023"), "06/11/2023")

    def test_digits_only_are_read_as_a_date(self):
        self.assertEqual(excel_filter.read_typed_date("06112023"), date(2023, 11, 6))

    def test_a_bare_day_means_that_day_of_the_report_month(self):
        self.assertEqual(excel_filter.read_typed_date("6", (2023, 11)), date(2023, 11, 6))
        self.assertEqual(excel_filter.read_typed_date("15", (2023, 11)), date(2023, 11, 15))

    def test_a_day_the_month_does_not_have_is_not_a_date(self):
        self.assertIsNone(excel_filter.read_typed_date("31", (2023, 11)))

    def test_other_separators_still_work(self):
        for text in ("06/11/2023", "06-11-2023", "06.11.2023", "2023-11-06", "06/11/23"):
            self.assertEqual(excel_filter.read_typed_date(text), date(2023, 11, 6), text)

    def test_nonsense_is_not_a_date(self):
        self.assertIsNone(excel_filter.read_typed_date("ontem", (2023, 11)))


class OneDay(unittest.TestCase):
    """With a single day in the filter, the chart is per hour, not one fat bar."""

    HEADERS = ["Data", "Horário", "Plan", "Produzido", "Principais perdas"]
    ROWS = [(2, [date(2023, 11, 6), "06h - 07h", 1885, 1700, ""]),
            (3, [date(2023, 11, 6), "07h - 08h", 1885, 1900, ""]),
            (4, [date(2023, 11, 6), "08h - 09h", 1885, 0, ""])]

    def test_hours_keep_the_shift_order_of_the_sheet(self):
        board = production_summary(self.HEADERS, self.ROWS)
        self.assertEqual([hour for hour, *_ in board["horas"]],
                         ["06h - 07h", "07h - 08h", "08h - 09h"])
        self.assertEqual(board["horas"][1], ("07h - 08h", 1885, 1900))

    def test_the_day_total_is_the_sum_of_its_hours(self):
        board = production_summary(self.HEADERS, self.ROWS)
        self.assertEqual(board["dias"], [(date(2023, 11, 6), 5655, 3600)])
        self.assertEqual(sum(p for _, p, _ in board["horas"]), board["plan"])

    def test_a_single_bar_does_not_stretch_across_the_chart(self):
        tk = __import__("tkinter")
        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        root.withdraw()
        canvas = tk.Canvas(root, width=900, height=200)
        canvas.pack()
        root.update()
        excel_filter.draw_combo_chart(canvas, [(date(2023, 11, 6), 30, 1885, 1700)])
        bar = next(i for i in canvas.find_withtag("slot") if canvas.type(i) == "rectangle")
        left, _, right, _ = canvas.coords(bar)
        self.assertLessEqual(right - left, 40)
        root.destroy()


class PointerLookup(unittest.TestCase):
    """Finding which bar the mouse is over, by point."""

    def setUp(self):
        tk = __import__("tkinter")
        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        self.root.withdraw()
        self.canvas = tk.Canvas(self.root, width=400, height=200)

    def tearDown(self):
        self.root.destroy()

    def test_finds_the_slot_of_the_day_under_the_point(self):
        days = [(date(2023, 11, d), 10, 1000, 800) for d in range(1, 6)]
        excel_filter.draw_combo_chart(self.canvas, days)
        left = [min(self.canvas.bbox(i)[0] for i in self.canvas.find_withtag(f"slot{i}"))
                for i in range(5)]
        for i, x in enumerate(left):
            self.assertEqual(excel_filter.tag_index(self.canvas, x + 4, 60, "slot"), i)

    def test_finds_the_machine_row_under_the_point(self):
        excel_filter.draw_rank_chart(self.canvas, [("Tampadora", 270), ("Gasser", 182)])
        row = self.canvas.bbox(*self.canvas.find_withtag("bar1"))
        middle = (row[1] + row[3]) / 2
        self.assertEqual(excel_filter.tag_index(self.canvas, 200, middle, "bar"), 1)

    def test_an_empty_chart_has_nothing_under_the_point(self):
        excel_filter.draw_combo_chart(self.canvas, [])
        self.assertIsNone(excel_filter.tag_index(self.canvas, 50, 50, "slot"))


class Slots(unittest.TestCase):
    """A short period is plotted hour by hour across the days."""

    HEADERS = ["Data", "Horário", "Plan", "Produzido"]
    ROWS = [(2, [date(2023, 11, 6), "06h - 07h", 1885, 1700]),
            (3, [date(2023, 11, 6), "07h - 08h", 1885, 1900]),
            (4, [date(2023, 11, 7), "06h - 07h", 1885, 1000]),
            (5, [date(2023, 11, 7), "07h - 08h", 1885, 1500])]

    def test_one_entry_per_day_and_hour_in_sheet_order(self):
        board = production_summary(self.HEADERS, self.ROWS)
        self.assertEqual([slot for slot, *_ in board["slots"]],
                         [(date(2023, 11, 6), "06h - 07h"), (date(2023, 11, 6), "07h - 08h"),
                          (date(2023, 11, 7), "06h - 07h"), (date(2023, 11, 7), "07h - 08h")])
        self.assertEqual(board["slots"][2][1:], (1885, 1000))

    def test_the_slots_add_up_to_the_days(self):
        board = production_summary(self.HEADERS, self.ROWS)
        self.assertEqual(sum(p for _, p, _ in board["slots"]), sum(p for _, p, _ in board["dias"]))


class SkuAndChangeover(unittest.TestCase):
    """SKU per day, with the changeover and the hour it happened."""

    HEADERS = ["Data", "SKU", "Trocas", "Período", "Principais perdas"]
    ROWS = [
        (2, [date(2023, 11, 16),
             "64834706 - Rexona Men Clinical Sport 150ml\n64830874 - Rexona Men Active Dry 150ml",
             "31m19s", "06h - 07h", "Change over: Troca de formato do Clinical Sport para Active Dry."]),
        (3, [date(2023, 11, 16), "64834706 - Rexona Men Clinical Sport 150ml\n64830874 - Rexona Men Active Dry 150ml",
             "31m19s", "07h - 08h", "Change over: Troca de formato do Clinical Sport para Active Dry."]),
        (4, [date(2023, 11, 17), "64830874 - Rexona Men Active Dry 150ml", "", "06h - 07h",
             "Tampadora - alarme de tampa."]),
    ]

    def by_day(self):
        return {day["day"]: (day["skus"], day["swaps"])
                for day in excel_filter.skus_by_day(self.HEADERS, self.ROWS)}

    def test_a_changeover_day_lists_both_skus_in_order(self):
        skus, _ = self.by_day()[date(2023, 11, 16)]
        self.assertEqual([code for code, _ in skus], ["64834706", "64830874"])
        self.assertIn("Clinical Sport", skus[0][1])

    def test_the_same_changeover_is_one_entry_with_every_hour_it_spans(self):
        _, swaps = self.by_day()[date(2023, 11, 16)]
        self.assertEqual(len(swaps), 1)
        hours, said, took = swaps[0]
        self.assertEqual(hours, ["06h - 07h", "07h - 08h"])
        self.assertEqual(took, "31m19s")  # vem do cabeçalho do turno, TROCA 1 - TEMPO
        self.assertIn("Clinical Sport para Active Dry", said)

    def test_change_over_written_with_a_space_still_counts(self):
        self.assertTrue(excel_filter.CHANGEOVER_RE.search("Change over: troca de formato"))
        self.assertTrue(excel_filter.CHANGEOVER_RE.search("changeover da linha"))
        self.assertTrue(excel_filter.CHANGEOVER_RE.search("Troca de variante do Dove"))

    def test_a_day_without_changeover_has_only_its_sku(self):
        skus, swaps = self.by_day()[date(2023, 11, 17)]
        self.assertEqual(([code for code, _ in skus], swaps), (["64830874"], []))

    def test_days_come_in_date_order(self):
        self.assertEqual([d["day"] for d in excel_filter.skus_by_day(self.HEADERS, self.ROWS)],
                         [date(2023, 11, 16), date(2023, 11, 17)])


class TimedChangeover(unittest.TestCase):
    """A changeover the header timed but nobody wrote down must not vanish."""

    HEADERS = ["Data", "Turno", "SKU", "Trocas", "Período", "Principais perdas"]

    def swaps(self, rows):
        return {d["day"]: d["swaps"] for d in excel_filter.skus_by_day(self.HEADERS, rows)}

    def test_timed_without_a_report_shows_up_with_no_hour(self):
        rows = [(2, [date(2023, 11, 8), 1, "64830874 - Rexona", "22min", "06h - 07h",
                     "Tampadora - alarme."])]
        swaps = self.swaps(rows)[date(2023, 11, 8)]
        self.assertEqual(len(swaps), 1)
        hours, said, took = swaps[0]
        self.assertEqual((hours, took), ([], "22min"))  # sem hora, e ela não é inventada
        self.assertIn("sem relato escrito", said)

    def test_a_shift_that_has_both_keeps_only_the_written_one(self):
        rows = [(2, [date(2023, 11, 16), 1, "64834706 - Clinical", "31m19s", "07h - 08h",
                     "Change over: troca de formato."])]
        swaps = self.swaps(rows)[date(2023, 11, 16)]
        self.assertEqual([(h, t) for h, _, t in swaps], [(["07h - 08h"], "31m19s")])

    def test_each_shift_is_judged_on_its_own(self):
        rows = [(2, [date(2023, 11, 9), 1, "64830874 - Rexona", "10min", "06h - 07h",
                     "Change over: troca."]),
                (3, [date(2023, 11, 9), 2, "64830874 - Rexona", "20min", "14h - 15h",
                     "Tampadora - alarme."])]
        swaps = self.swaps(rows)[date(2023, 11, 9)]
        self.assertEqual(sorted((tuple(h), t) for h, _, t in swaps),
                         [((), "20min"), (("06h - 07h",), "10min")])


class ChangeoverLines(unittest.TestCase):
    """What counts as a changeover and what is only talk about one."""

    def is_event(self, line):
        return bool(excel_filter.CHANGEOVER_RE.match(line))

    def test_the_headline_of_the_event_counts(self):
        for line in ("Change over: Troca de formato do Clinical Sport para Active Dry",
                     "CHANGE OVER DO DOVE AMORA 12X150ML PARA REXONA W INVISIBLE",
                     "Changeover realizado no início do turno",
                     "Troca de variante do Dove Lavanda para Dove Original",
                     "* Troca de formato para Tripack"):
            self.assertTrue(self.is_event(line), line)

    def test_the_typos_in_the_sheet_still_count(self):
        # escrito assim no relatório: "CHANGER OVER", "Chang Over"
        self.assertTrue(self.is_event("CHANGER OVER - Troca de formato e variante"))
        self.assertTrue(self.is_event("Chang Over: Troca de variante de Rexona W Invisible"))

    def test_talking_about_the_changeover_is_not_another_changeover(self):
        for line in ("Durante a troca deu alarme nivel 2 na Novopack, atrasando o change over",
                     "Seladora de Caixas - Durante o changeover foi ajustado a altura do encolado",
                     "Atraso na liberação da linha devido a falha no change over",
                     "Checklist operação da troca de variante OK",
                     "Tampadora - Após o change over excesso de rejeito de lata"):
            self.assertFalse(self.is_event(line), line)

    def test_the_context_line_stops_being_a_second_changeover(self):
        headers = ["Data", "Turno", "SKU", "Trocas", "Período", "Principais perdas"]
        rows = [(2, [date(2023, 11, 16), 1, "64834706 - Clinical Sport", "31m19s", "06h - 07h",
                     "Change over: Troca de formato do Clinical Sport para Active Dry\n"
                     "Durante a troca deu alarme nivel 2 na Novopack, atrasando o change over"])]
        swaps = excel_filter.skus_by_day(headers, rows)[0]["swaps"]
        self.assertEqual(len(swaps), 1)
        self.assertIn("Troca de formato", swaps[0][1])


class BundleMonths(unittest.TestCase):
    """Which month each workbook of a zip belongs to."""

    def test_the_month_in_the_file_name(self):
        self.assertEqual(excel_filter.month_from_name("Relatório - Linha 1- V8- Março.xlsx"),
                         (None, 3))
        self.assertEqual(excel_filter.month_from_name("Turno 2023-11.xlsx"), (2023, 11))
        self.assertEqual(excel_filter.month_from_name("turno_11-2023.xlsx"), (2023, 11))
        self.assertIsNone(excel_filter.month_from_name("Relatório de Turno - Linha 2 V8.xlsx"))

    def test_a_month_that_is_not_a_month_is_not_read(self):
        self.assertIsNone(excel_filter.month_from_name("turno 2023-13.xlsx"))

    def zipped(self, files):
        path = Path(tempfile.mkdtemp()) / "pacote.zip"
        with zipfile.ZipFile(path, "w") as bundle:
            for name, book in files:
                where = Path(tempfile.mkdtemp()) / name
                book.save(where)
                bundle.write(where, f"pacote/{name}")
        return path

    def report(self, month=None):
        """A tiny shift report, optionally with a planning sheet naming a month."""
        book = Workbook()
        book.remove(book.active)
        if month:
            plan = book.create_sheet("Plan.Prod")
            for i in range(28):
                plan.cell(1 + i, 2).value = datetime(month[0], month[1], 1 + i)
        for day in range(1, 6):
            sheet = book.create_sheet(str(day))
            sheet.cell(26, 2).value = "Horário"
            sheet.cell(26, 14).value = "Principais Perdas"
            for h in range(8):
                sheet.cell(27 + h, 2).value = f"{6 + h:02d}h - {7 + h:02d}h"
        return book

    def test_every_workbook_of_the_zip_is_read_with_its_own_month(self):
        path = self.zipped([("Relatório - Janeiro.xlsx", self.report((2026, 1))),
                            ("Relatório - Abril.xlsx", self.report((2026, 4)))])
        data = load(path)
        self.assertEqual(sorted(when for _, when, _ in data["months"]), [(2026, 1), (2026, 4)])
        self.assertEqual(len(data["rows"]), 2 * 5 * 8)  # 2 arquivos x 5 dias x 8 horas

    def test_the_file_name_wins_when_the_planning_sheet_disagrees(self):
        path = self.zipped([("Relatório - Janeiro.xlsx", self.report((2026, 1))),
                            ("Relatório - Fevereiro.xlsx", self.report((2023, 11)))])
        found = {name: (when, source) for name, when, source in load(path)["months"]}
        when, source = found["Relatório - Fevereiro.xlsx"]
        self.assertEqual(when, (2026, 2))  # mês do nome, ano emprestado do outro arquivo
        self.assertIn("dizia 11/2023", source)

    def test_a_zip_without_spreadsheets_says_so(self):
        path = Path(tempfile.mkdtemp()) / "vazio.zip"
        with zipfile.ZipFile(path, "w") as bundle:
            bundle.writestr("leia.txt", "nada aqui")
        with self.assertRaises(ValueError):
            load(path)

    def test_the_rows_of_a_borrowed_year_are_kept_and_dated_right(self):
        path = self.zipped([("Relatório - Janeiro.xlsx", self.report((2026, 1))),
                            ("Relatório - Fevereiro.xlsx", self.report((2023, 11)))])
        data = load(path)
        self.assertEqual(len(data["rows"]), 2 * 5 * 8)  # nada se perdeu no caminho
        date_at = data["headers"].index("Data")
        fevereiro = {r[date_at] for _, r in data["rows"] if r[date_at].month == 2}
        self.assertEqual(min(fevereiro), date(2026, 2, 1))
        self.assertEqual(max(fevereiro), date(2026, 2, 5))

    def test_the_headers_survive_a_workbook_that_could_not_be_read(self):
        book = Workbook()  # sem abas de dia: não é relatório de turno
        book.active["A1"] = "qualquer coisa"
        path = self.zipped([("Relatório - Janeiro.xlsx", self.report((2026, 1))),
                            ("outra coisa.xlsx", book)])
        data = load(path)
        self.assertIn("Data", data["headers"])
        self.assertEqual(len(data["rows"]), 5 * 8)
        self.assertIn("não é relatório de turno",
                      [source for name, _, source in data["months"] if name == "outra coisa.xlsx"])


class MonthlySeries(unittest.TestCase):
    """More than a month of bars is a picket fence: the chart goes monthly."""

    def rows(self, days):
        return [(i, [day, 1885 * 24, 1700 * 24]) for i, day in enumerate(days, 2)]

    HEADERS = ["Data", "Plan", "Produzido"]

    def test_one_entry_per_month_with_the_month_total(self):
        days = [date(2026, 1, d) for d in range(1, 32)] + [date(2026, 2, d) for d in range(1, 29)]
        board = production_summary(self.HEADERS, self.rows(days))
        self.assertEqual([month for month, *_ in board["meses"]],
                         [date(2026, 1, 1), date(2026, 2, 1)])
        self.assertEqual(board["meses"][0][1], 1885 * 24 * 31)  # janeiro inteiro somado
        self.assertEqual(sum(p for _, p, _ in board["meses"]), board["plan"])

    def test_the_chart_labels_a_monthly_series_by_month(self):
        tk = __import__("tkinter")
        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        root.withdraw()
        canvas = tk.Canvas(root, width=600, height=200)
        excel_filter.draw_combo_chart(canvas, [(date(2026, 1, 1), 5, 100, 90),
                                               (date(2026, 2, 1), 5, 100, 110)], monthly=True)
        labels = [canvas.itemcget(i, "text") for i in canvas.find_all()
                  if canvas.type(i) == "text"]
        self.assertIn("01/26", labels)
        self.assertIn("02/26", labels)
        root.destroy()

    def test_a_daily_series_keeps_the_day_numbers(self):
        tk = __import__("tkinter")
        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        root.withdraw()
        canvas = tk.Canvas(root, width=600, height=200)
        excel_filter.draw_combo_chart(canvas, [(date(2026, 1, 5), 5, 100, 90),
                                               (date(2026, 1, 6), 5, 100, 110)])
        labels = [canvas.itemcget(i, "text") for i in canvas.find_all()
                  if canvas.type(i) == "text"]
        self.assertIn("05", labels)
        root.destroy()


class Merging(unittest.TestCase):
    """Two bundles seen as one."""

    def dataset(self, days, line="D01", headers=("Data", "Linha", "Plan", "Produzido")):
        rows = [(i, [day, line, 1885, 1700]) for i, day in enumerate(days, 2)]
        return {"kind": "turno", "sheet": "x", "headers": list(headers), "rows": rows,
                "month": (2026, 1), "months": [(f"{line}.zip", (2026, 1), "aba Plan.Prod")],
                "path": f"/tmp/{line}.zip"}

    def test_the_rows_of_both_come_together_renumbered(self):
        joined = excel_filter.merge([self.dataset([date(2026, 1, 1)], "D01"),
                                     self.dataset([date(2026, 1, 1), date(2026, 1, 2)], "D02")])
        self.assertEqual(len(joined["rows"]), 3)
        self.assertEqual([n for n, _ in joined["rows"]], [1, 2, 3])
        self.assertEqual(sorted({r[1] for _, r in joined["rows"]}), ["D01", "D02"])

    def test_it_keeps_where_each_one_came_from(self):
        joined = excel_filter.merge([self.dataset([date(2026, 1, 1)], "D01"),
                                     self.dataset([date(2026, 1, 1)], "D02")])
        self.assertEqual(joined["sources"], ["/tmp/D01.zip", "/tmp/D02.zip"])
        self.assertEqual(len(joined["months"]), 2)

    def test_adding_a_third_keeps_the_first_two(self):
        um, dois, tres = (self.dataset([date(2026, 1, 1)], linha) for linha in ("D01", "D02", "D03"))
        juntos = excel_filter.merge([excel_filter.merge([um, dois]), tres])
        self.assertEqual(juntos["sources"], ["/tmp/D01.zip", "/tmp/D02.zip", "/tmp/D03.zip"])
        self.assertEqual(sorted({r[1] for _, r in juntos["rows"]}), ["D01", "D02", "D03"])
        self.assertEqual(len(juntos["months"]), 3)
        self.assertEqual(juntos["sheet"], "3 arquivos")

    def test_reports_of_different_shapes_are_not_merged(self):
        with self.assertRaises(ValueError):
            excel_filter.merge([self.dataset([date(2026, 1, 1)]),
                                self.dataset([date(2026, 1, 1)],
                                             headers=("Data", "Outra coisa"))])

    def test_merging_nothing_is_an_error_not_an_empty_screen(self):
        with self.assertRaises(ValueError):
            excel_filter.merge([])


class ForgetHistory(unittest.TestCase):
    def setUp(self):
        excel_filter.HISTORY_FILE = Path(tempfile.mkdtemp()) / "historico.json"
        excel_filter.remember("/tmp/a.zip", {"rows": 10, "month": None})
        excel_filter.remember("/tmp/b.zip", {"rows": 20, "month": None})

    def test_forgetting_one_keeps_the_others(self):
        self.assertEqual([e["path"] for e in excel_filter.forget("/tmp/a.zip")], ["/tmp/b.zip"])
        self.assertEqual([e["path"] for e in excel_filter.history()], ["/tmp/b.zip"])

    def test_forgetting_all(self):
        self.assertEqual(excel_filter.forget_all(), [])
        self.assertEqual(excel_filter.history(), [])

    def test_forgetting_what_is_not_there_changes_nothing(self):
        self.assertEqual(len(excel_filter.forget("/tmp/nao_existe.zip")), 2)


class ChangeoverTime(unittest.TestCase):
    """The changeover duration as a number, and summed per day."""

    def test_the_spellings_the_shift_leaders_use(self):
        minutes = excel_filter.changeover_minutes
        self.assertAlmostEqual(minutes("31m19s"), 31 + 19 / 60)
        self.assertEqual(minutes("22min"), 22)
        self.assertEqual(minutes("45 minutos"), 45)
        self.assertEqual(minutes("1h05m"), 65)
        self.assertEqual(minutes("2h"), 120)
        self.assertEqual(minutes("0 Minutos (TROCA DE ORDEM)"), 0)

    def test_what_is_not_a_duration_is_not_a_zero(self):
        for text in ("N/A", "", None, "abc"):
            self.assertIsNone(excel_filter.changeover_minutes(text))

    def test_one_duration_per_shift_even_though_every_hour_repeats_it(self):
        headers = ["Data", "Turno", "Trocas"]
        rows = [(i, [date(2026, 1, 5), 1, "30min"]) for i in range(2, 10)]  # 8 horas do turno
        rows += [(i, [date(2026, 1, 5), 2, "10min"]) for i in range(10, 18)]
        self.assertEqual(excel_filter.changeover_time_by_day(headers, rows),
                         [(date(2026, 1, 5), 40)])  # 30 + 10, não 8x isso

    def test_days_without_changeover_count_zero_not_nothing(self):
        headers = ["Data", "Turno", "Trocas"]
        rows = [(2, [date(2026, 1, 5), 1, "30min"]), (3, [date(2026, 1, 6), 1, "N/A"])]
        self.assertEqual(excel_filter.changeover_time_by_day(headers, rows),
                         [(date(2026, 1, 5), 30), (date(2026, 1, 6), 0)])


class SkuProgression(unittest.TestCase):
    """Each changeover moves the day on to the next SKU of the cell."""

    HEADERS = ["Data", "Turno", "SKU", "Trocas", "Período", "Principais perdas", "Arquivo"]

    def day(self, rows):
        return excel_filter.skus_by_day(self.HEADERS, rows)[0]

    def test_the_changeovers_come_in_the_order_they_happened(self):
        sku = "1111111 - Primeiro\n2222222 - Segundo\n3333333 - Terceiro"
        rows = [(2, [date(2026, 1, 12), 2, sku, "20min", "14h - 15h",
                     "Change over: troca para o terceiro", "janeiro.xlsx"]),
                (3, [date(2026, 1, 12), 1, sku, "31m19s", "08h - 09h",
                     "Change over: troca para o segundo", "janeiro.xlsx"])]
        day = self.day(rows)
        self.assertEqual([swap[0] for swap in day["swaps"]], [["08h - 09h"], ["14h - 15h"]])
        self.assertEqual([code for code, _ in day["skus"]], ["1111111", "2222222", "3333333"])

    def test_the_day_says_which_report_it_came_from(self):
        rows = [(2, [date(2026, 1, 12), 1, "1111111 - Um", "", "06h - 07h", "", "janeiro.xlsx"])]
        self.assertEqual(self.day(rows)["file"], "janeiro.xlsx")

    def test_a_changeover_with_no_hour_goes_last(self):
        rows = [(2, [date(2026, 1, 12), 1, "1111111 - Um\n2222222 - Dois", "20min", "08h - 09h",
                     "Change over: troca", "x.xlsx"]),
                (3, [date(2026, 1, 12), 2, "1111111 - Um\n2222222 - Dois", "15min", "14h - 15h",
                     "Tampadora - alarme", "x.xlsx"])]
        hours = [swap[0] for swap in self.day(rows)["swaps"]]
        self.assertEqual(hours, [["08h - 09h"], []])  # o cronometrado sem relato vai por último


class UnknownPlan(unittest.TestCase):
    """An hour whose Plan cell is empty must not pass for a beaten target."""

    HEADERS = ["Data", "Horário", "Plan", "Produzido"]

    def test_a_bucket_with_no_readable_plan_reports_none(self):
        rows = [(2, [date(2026, 2, 3), "06h - 07h", None, 1500])]
        board = production_summary(self.HEADERS, rows)
        self.assertEqual(board["dias"], [(date(2026, 2, 3), None, 1500)])
        self.assertEqual(board["horas"], [("06h - 07h", None, 1500)])

    def test_the_hours_that_do_have_a_plan_still_add_up(self):
        rows = [(2, [date(2026, 2, 3), "06h - 07h", None, 1500]),
                (3, [date(2026, 2, 3), "07h - 08h", 1925, 1800])]
        self.assertEqual(production_summary(self.HEADERS, rows)["dias"],
                         [(date(2026, 2, 3), 1925, 3300)])  # o plano ilegível não vira zero

    def test_a_month_of_unknown_plans_stays_unknown(self):
        rows = [(i, [date(2026, 2, d), "06h - 07h", None, 1000]) for i, d in enumerate((3, 4), 2)]
        self.assertEqual(production_summary(self.HEADERS, rows)["meses"],
                         [(date(2026, 2, 1), None, 2000)])

    def test_the_bar_is_neither_green_nor_red_when_there_is_no_target(self):
        tk = __import__("tkinter")
        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"sem display: {error}")
        root.withdraw()
        canvas = tk.Canvas(root, width=400, height=200)
        excel_filter.draw_combo_chart(canvas, [(date(2026, 2, 3), 0, None, 1500)])
        painted = [canvas.itemcget(i, "fill") for i in canvas.find_withtag("slot")
                   if canvas.type(i) == "rectangle" and canvas.itemcget(i, "fill")]
        # cinza-ardósia: sem plano não há "bateu"
        self.assertEqual(painted, [excel_filter.BOARD_COLOURS["unknown"]])
        ticks = [i for i in canvas.find_withtag("slot") if canvas.type(i) == "line"]
        self.assertEqual(ticks, [])  # e sem marca de plano, porque não há alvo
        root.destroy()

    def test_the_format_of_the_hour_is_carried_as_data(self):
        # Hagrid, Golias, Axe, Tripack: o app não tem regra por formato, só transporta
        headers = ["Data", "Formato", "Plan", "Produzido", "Principais perdas"]
        rows = [(2, [date(2026, 1, 5), "Hagrid", 1482, 1400, "Gasser - alarme."]),
                (3, [date(2026, 1, 5), "Golias", 1314, 1300, "Tampadora - alarme."])]
        result = summarize(headers, rows, date_col="Data", text_col="Principais perdas")
        self.assertEqual(column(result, "Formato"), ["Hagrid", "Golias"])
        self.assertEqual(result["totals"]["Plan"], 1482 + 1314)


class FastReader(unittest.TestCase):
    """The streaming reader must agree with the slow one, cell for cell."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        book = Workbook()
        book.remove(book.active)
        plan = book.create_sheet("Plan.Prod")
        for i in range(30):
            plan.cell(1 + i, 2).value = datetime(2026, 1, 1 + i)
        for day in range(1, 9):
            sheet = book.create_sheet(str(day))
            for shift, base in enumerate((2, 35, 68), 1):
                if shift == 1:
                    sheet.cell(2, base).value = f"LINHA D01 - {shift}º TURNO"
                sheet.cell(6, base).value = "SKU DESCRIÇÃO"
                sheet.cell(6, base + 3).value = "1111111 - Um\n2222222 - Dois"
                sheet.cell(6, base + 13).value = "TROCA 1 - TEMPO"
                sheet.cell(6, base + 16).value = "31m19s"
                for label, off in (("Horário", 0), ("Formato", 2), ("Plan", 4),
                                   ("Produzido", 6), ("Apontado", 8), ("Principais Perdas", 12)):
                    sheet.cell(26, base + off).value = label
                for h in range(8):
                    row = 27 + h
                    sheet.cell(row, base).value = f"{(6 * shift + h) % 24:02d}h - {(6 * shift + h + 1) % 24:02d}h"
                    sheet.cell(row, base + 2).value = "Hagrid"
                    sheet.cell(row, base + 4).value = 1482
                    sheet.cell(row, base + 6).value = 1400 + h
                sheet.cell(27, base + 12).value = "Change over: troca de formato.\nGasser - alarme."
                sheet.merge_cells(start_row=27, end_row=30, start_column=base + 12,
                                  end_column=base + 14)
        self.path = Path(self.tmp) / "turno.xlsx"
        book.save(self.path)

    def read_both(self):
        slow = load_workbook(self.path, data_only=True)
        by_hand = excel_filter.read_shift_report(slow, 2026, 1)
        fast_book = load_workbook(self.path, data_only=True, read_only=True)
        days = [n for n in fast_book.sheetnames if n.isdigit()]
        streamed = excel_filter.read_shift_report(
            fast_book, 2026, 1, sheets=excel_filter.fast_sheets(self.path, days))
        return by_hand, streamed

    def test_both_readers_give_the_same_headers_and_rows(self):
        (headers, rows), (fast_headers, fast_rows) = self.read_both()
        self.assertEqual(headers, fast_headers)
        self.assertEqual(rows, fast_rows)
        self.assertEqual(len(rows), 8 * 3 * 8)

    def test_the_streamed_sheet_keeps_the_merged_ranges(self):
        fast = excel_filter.fast_sheets(self.path, ["1"])["1"]
        spans = [(r.min_row, r.max_row, r.min_col, r.max_col) for r in fast.merged_cells.ranges]
        self.assertIn((27, 30, 14, 16), spans)  # a mescla do 1º turno

    def test_the_streamed_sheet_answers_by_coordinate(self):
        fast = excel_filter.fast_sheets(self.path, ["1"])["1"]
        self.assertEqual(fast.cell(2, 2).value, "LINHA D01 - 1º TURNO")
        self.assertEqual(fast.cell(26, 2).value, "Horário")
        self.assertIsNone(fast.cell(60, 60).value)  # célula vazia não estoura

    def test_the_period_and_the_changeover_survive_the_fast_path(self):
        _, (headers, rows) = self.read_both()
        first = rows[0][1]
        self.assertEqual(first[headers.index("Período")], "06h - 10h")  # a mescla de 4 horas
        self.assertEqual(first[headers.index("Trocas")], "31m19s")
        self.assertEqual(first[headers.index("Formato")], "Hagrid")


class ChartColumns(unittest.TestCase):
    """Changeover and production share one axis, one column per day or month."""

    def numbers(self, days):
        return {"dias": [(d, 1000, 900) for d in days], "meses": [],
                "trocas_por_dia": [(d, 10) for d in days]}

    def test_each_day_carries_its_changeover_and_production(self):
        days = [date(2026, 1, 5), date(2026, 1, 6)]
        columns, monthly = excel_filter.chart_columns(self.numbers(days))
        self.assertFalse(monthly)
        self.assertEqual(columns, [(date(2026, 1, 5), 10, 1000, 900), (date(2026, 1, 6), 10, 1000, 900)])

    def test_more_than_a_month_goes_monthly_and_adds_the_changeover_up(self):
        days = [date(2026, 1, d) for d in range(1, 32)] + [date(2026, 2, 1)]
        numbers = self.numbers(days)
        numbers["meses"] = [(date(2026, 1, 1), 31000, 27900), (date(2026, 2, 1), 1000, 900)]
        columns, monthly = excel_filter.chart_columns(numbers)
        self.assertTrue(monthly)
        self.assertEqual([(when, minutes) for when, minutes, _, _ in columns],
                         [(date(2026, 1, 1), 310), (date(2026, 2, 1), 10)])



class Access(unittest.TestCase):
    """Who is the admin. Test credentials only: the real ones never enter the repo."""

    def setUp(self):
        self.path = Path(tempfile.mkdtemp()) / "admin.json"
        excel_filter.make_admin("admin.teste", "Senha-Certa1", self.path)

    def test_the_right_login_and_password_open_the_admin(self):
        self.assertTrue(excel_filter.is_admin("admin.teste", "Senha-Certa1", self.path))
        self.assertTrue(excel_filter.is_admin("  Admin.Teste ", "Senha-Certa1", self.path))

    def test_anything_else_stays_a_viewer(self):
        for user, password in (("admin.teste", "senha-certa1"), ("outra", "Senha-Certa1"),
                               ("", ""), ("admin.teste", "")):
            self.assertFalse(excel_filter.is_admin(user, password, self.path), (user, password))

    def test_the_password_itself_is_never_written(self):
        self.assertNotIn("Senha-Certa1", self.path.read_text())

    def test_no_file_or_a_broken_one_means_no_admin(self):
        self.assertFalse(excel_filter.is_admin("admin.teste", "Senha-Certa1",
                                               self.path.with_name("nada.json")))
        for broken in ("[]", "7", "{}", "não é json", '{"user": "admin.teste", "salt": "zz"}'):
            self.path.write_text(broken)
            self.assertFalse(excel_filter.is_admin("admin.teste", "Senha-Certa1", self.path), broken)


class SavedBase(unittest.TestCase):
    """What the admin feeds stays on disk, once per line and month."""

    def setUp(self):
        self.folder = Path(tempfile.mkdtemp()) / "base"

    def report(self, line, month):
        book = Workbook()
        book.remove(book.active)
        plan = book.create_sheet("Plan.Prod")
        for i in range(28):
            plan.cell(1 + i, 2).value = datetime(month[0], month[1], 1 + i)
        for day in range(1, 6):
            sheet = book.create_sheet(str(day))
            sheet.cell(1, 2).value = f"LINHA {line} - 1º TURNO"
            sheet.cell(26, 2).value = "Horário"
            sheet.cell(26, 14).value = "Principais Perdas"
            for h in range(8):
                sheet.cell(27 + h, 2).value = f"{6 + h:02d}h - {7 + h:02d}h"
            sheet.cell(27, 14).value = f"Gasser - alarme no dia {day}"
        return book

    def bundle(self, months, name="pacote.zip", line="D01"):
        path = Path(tempfile.mkdtemp()) / name
        with zipfile.ZipFile(path, "w") as package:
            for month in months:
                where = Path(tempfile.mkdtemp()) / f"Linha 1 - {month[1]:02d}.xlsx"
                self.report(line, month).save(where)
                package.write(where, where.name)
        return load(path)

    def test_what_goes_in_comes_back_the_same_with_dates_as_dates(self):
        data = self.bundle([(2026, 1), (2026, 2)])
        saved, replaced = excel_filter.base_add(data, "pacote.zip", self.folder)
        self.assertEqual((len(saved), replaced), (2, []))
        back = excel_filter.base_load(self.folder)
        self.assertEqual(sorted(map(repr, (v for _, v in back["rows"]))),
                         sorted(map(repr, (v for _, v in data["rows"]))))
        day = back["rows"][0][1][back["headers"].index("Data")]
        self.assertIs(type(day), date)
        self.assertEqual([when for _, when, _ in back["months"]], [(2026, 1), (2026, 2)])

    def test_sending_the_same_month_again_replaces_it_instead_of_doubling(self):
        data = self.bundle([(2026, 1), (2026, 2)])
        excel_filter.base_add(data, "pacote.zip", self.folder)
        files_before = sorted(p.name for p in self.folder.iterdir())
        again = self.bundle([(2026, 1), (2026, 2)], name="pacote (1).zip")
        _, replaced = excel_filter.base_add(again, "pacote (1).zip", self.folder)
        self.assertEqual(len(replaced), 2)
        self.assertEqual(len(excel_filter.base_entries(self.folder)), 2)
        self.assertEqual(len(excel_filter.base_load(self.folder)["rows"]), len(data["rows"]))
        self.assertEqual(len(list(self.folder.iterdir())), len(files_before))  # old files went away

    def test_a_new_month_or_another_line_adds_to_what_is_there(self):
        excel_filter.base_add(self.bundle([(2026, 1)]), "d01.zip", self.folder)
        excel_filter.base_add(self.bundle([(2026, 2)]), "d01.zip", self.folder)
        excel_filter.base_add(self.bundle([(2026, 1)], line="D02"), "d02.zip", self.folder)
        keys = sorted(e["key"] for e in excel_filter.base_entries(self.folder))
        self.assertEqual(keys, ["D01|2026-01", "D01|2026-02", "D02|2026-01"])

    def test_deleting_takes_the_rows_and_the_file_away(self):
        excel_filter.base_add(self.bundle([(2026, 1), (2026, 2)]), "pacote.zip", self.folder)
        january = next(e for e in excel_filter.base_entries(self.folder) if e["month"] == [2026, 1])
        left = excel_filter.base_remove([january["id"]], self.folder)
        self.assertEqual([e["month"] for e in left], [[2026, 2]])
        self.assertFalse((self.folder / f"{january['id']}.json").exists())
        back = excel_filter.base_load(self.folder)
        self.assertEqual({when for _, when, _ in back["months"]}, {(2026, 2)})
        excel_filter.base_remove([left[0]["id"]], self.folder)
        self.assertIsNone(excel_filter.base_load(self.folder))

    def test_an_empty_or_broken_base_is_an_empty_base(self):
        self.assertIsNone(excel_filter.base_load(self.folder))
        self.folder.mkdir(parents=True)
        for broken in ("[]", "7", "não é json"):
            (self.folder / "indice.json").write_text(broken)
            self.assertIsNone(excel_filter.base_load(self.folder))

    def test_an_edited_index_cannot_reach_outside_the_folder(self):
        self.folder.mkdir(parents=True)
        outside = self.folder.parent / "fora.json"
        outside.write_text("{}")
        (self.folder / "indice.json").write_text(json.dumps([{"id": "../fora", "key": "x|?"}]))
        self.assertEqual(excel_filter.base_entries(self.folder), [])
        excel_filter.base_remove(["../fora"], self.folder)
        self.assertTrue(outside.exists())

    def test_a_lost_file_is_reported_not_hidden(self):
        excel_filter.base_add(self.bundle([(2026, 1), (2026, 2)]), "pacote.zip", self.folder)
        lost = excel_filter.base_entries(self.folder)[0]
        (self.folder / f"{lost['id']}.json").unlink()
        self.assertEqual(excel_filter.base_load(self.folder)["missing"], [lost["file"]])

    def test_only_shift_reports_with_a_read_month_go_in(self):
        with self.assertRaises(ValueError):
            excel_filter.base_add({"kind": "tabela", "headers": [], "rows": []}, "x", self.folder)
        with self.assertRaises(ValueError):
            excel_filter.base_add({"kind": "turno", "month_guessed": True, "headers": [], "rows": []},
                                  "x", self.folder)


class Links(unittest.TestCase):
    """Feeding the base from a link: the file, or a clear no."""

    def serve(self, status, body, headers=()):
        import http.server
        import threading

        class Answer(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                self.send_response(status)
                for name, value in headers:
                    self.send_header(name, value)
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *args):
                pass

        server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Answer)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        return f"http://127.0.0.1:{server.server_port}/arquivo"

    def test_a_sharepoint_link_asks_for_the_file_itself(self):
        self.assertEqual(excel_filter.direct_link("https://empresa.sharepoint.com/:u:/s/x/EaB?e=abc"),
                         "https://empresa.sharepoint.com/:u:/s/x/EaB?e=abc&download=1")
        self.assertEqual(excel_filter.direct_link("https://exemplo.com/a.zip"),
                         "https://exemplo.com/a.zip")
        with self.assertRaises(ValueError):
            excel_filter.direct_link("D01_Relatorio.zip")

    def test_a_zip_is_saved_under_the_name_the_server_gave(self):
        body = io.BytesIO()
        with zipfile.ZipFile(body, "w") as package:
            package.writestr("pacote/leia.txt", "oi")
        url = self.serve(200, body.getvalue(),
                         [("Content-Disposition", 'attachment; filename="D01_Relatorio.zip"')])
        path = excel_filter.fetch_link(url, tempfile.mkdtemp())
        self.assertEqual((path.name, path.read_bytes()), ("D01_Relatorio.zip", body.getvalue()))

    def test_a_workbook_comes_back_as_xlsx(self):
        where = Path(tempfile.mkdtemp()) / "x.xlsx"
        Workbook().save(where)
        path = excel_filter.fetch_link(self.serve(200, where.read_bytes()), tempfile.mkdtemp())
        self.assertEqual(path.suffix, ".xlsx")

    def test_a_login_page_is_said_never_read_as_a_spreadsheet(self):
        for status, body in ((200, b"<html>Entrar na sua conta</html>"), (403, b"proibido"),
                             (401, b"")):
            with self.assertRaises(ValueError) as said:
                excel_filter.fetch_link(self.serve(status, body), tempfile.mkdtemp())
            self.assertIn("login", str(said.exception))
