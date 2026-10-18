"""Shift-report summary: filter a spreadsheet by date range, keyword and shift.

Code is English, the UI is pt-BR on purpose. Run with: python excel_filter.py
"""

import base64
import io
import json
import queue
import calendar
import hashlib
import hmac
import posixpath
import re
import secrets
import shutil
import sys
import tempfile
import threading
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

BUILD = "18/10 02:42"  # carimbo do último commit, para saber qual versão está rodando
DATE_FORMATS = ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y",
                "%d%m%Y", "%d%m%y")
HOUR_RE = re.compile(r"\s*\d{1,2}\s*(h|:\d{2})\s*[-–às]*\s*(\d{1,2}\s*(h|:\d{2}))?\s*", re.I)
# shift -> (first hour, hour it ends); the third one crosses midnight
SHIFTS = {1: (6, 14), 2: (14, 22), 3: (22, 6)}
SHIFT_LABELS = {n: f"{n}º turno ({a:02d}h às {b:02d}h)" for n, (a, b) in SHIFTS.items()}
NS_MAIN ="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def parse_date(value):
    """Cell (date/datetime/text) -> date, or None when it cannot be read."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def shift_of(hour_text):
    """"06h - 07h" -> 1, "15:00" -> 2, "23h - 00h" -> 3; None when no hour can be read.

    Goes by the hour the slot STARTS at: "13h - 14h" is still the first shift.
    """
    found = re.search(r"\d{1,2}", str(hour_text or ""))
    if not found or int(found.group()) > 23:
        return None
    hour = int(found.group())
    for number, (first, until) in SHIFTS.items():
        if (first <= hour < until) if first < until else (hour >= first or hour < until):
            return number
    return None


def mask_date(text):
    """What was typed, as dd/mm/aaaa — the slashes appear on their own."""
    digits = re.sub(r"\D", "", text)[:8]
    return "/".join(part for part in (digits[:2], digits[2:4], digits[4:]) if part)


def read_typed_date(text, month=None):
    """Text from a date field -> date, or None.

    A bare day ("6", "15") means that day of the report's month, which is the
    common case here: the whole spreadsheet is one month.
    """
    text = text.strip()
    if not text:
        return None
    if month and re.fullmatch(r"\d{1,2}", text):
        year, number = month
        try:
            return date(year, number, int(text))
        except ValueError:
            return None
    return parse_date(text)


def _header(sheet):
    return [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]


def read(path, sheet=None, book=None):
    """(sheet name, headers, [(row number, values)]) with calculated values.

    Without `sheet`, picks the one that is an actual table: filled header on row 1
    and the most rows — "Tabela de Dados" among the 35 sheets of the shift report.
    Reads once: the UI keeps the rows and refilters in memory, otherwise every
    keystroke would reopen a file that is 59MB here.
    """
    book = book or load_workbook(path, data_only=True)
    if sheet is None:
        ranked = [(sum(1 for h in _header(book[n]) if h), book[n].max_row, n) for n in book.sheetnames]
        sheet = max(ranked, key=lambda c: (c[0] >= 3, c[1]))[2]
    ws = book[sheet]
    headers = _header(ws)
    rows = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None and str(v).strip() != "" for v in values):
            rows.append((r, values))
    return sheet, headers, rows


# --- Shift report (the "Relatório de Turno" workbook) --------------------------
# Its "Tabela de Dados" sheet is broken: column A was dragged down one day per
# ROW instead of per day, so it drifts 01/09 -> 24/10 for a single month, and the
# INDIRECT that pulls the text picks the day sheet by TEXT(A,"d") — meaning each
# block of 24 hours reads 24 DIFFERENT days. The day sheets are the real source.
LINE_RE = re.compile(r"LINHA\s+([A-Z]?\d+[A-Z]?)", re.I)  # "LINHA D02 - 1º TURNO" -> D02
SHIFT_FIELD_NAMES = {"horario": "horário", "formato": "formato", "plan": "plan",
                     "produzido": "produzido", "apontado": "apontado",
                     "perdas": "principais perdas"}
SHIFT_HOURS = 8  # each shift block lists eight hours
# where the screen finds each thing in a dataset built from the day sheets
SHIFT_COLUMNS = {"date": "Data", "hour": "Horas", "label": "Período", "text": "Principais perdas",
                 "line": "Linha"}


class _Cell:
    """What the reader asks of a cell: its value and where it is."""
    __slots__ = ("value", "row", "column")

    def __init__(self, value, row, column):
        self.value, self.row, self.column = value, row, column


class _Range:
    """A merged range, the four numbers the reader uses."""
    __slots__ = ("min_row", "max_row", "min_col", "max_col")

    def __init__(self, min_row, max_row, min_col, max_col):
        self.min_row, self.max_row = min_row, max_row
        self.min_col, self.max_col = min_col, max_col


class _Merged:
    def __init__(self, ranges):
        self.ranges = ranges


class FastSheet:
    """A day sheet read in streaming mode: only the rows the report lives in.

    openpyxl in normal mode takes 25s to open one 58MB workbook and 0.4s to pull
    the data out of it — it parses every cell of every sheet. In read-only mode the
    same file opens in 0.1s, but read-only sheets have no random access and no
    merged ranges, which the reader needs. So the top rows are copied into a grid
    and the merges are read straight from the sheet XML.
    """
    ROWS = 70  # header block, hour rows and some slack

    def __init__(self, values, merges):
        self._values, self.merged_cells = values, _Merged(merges)

    def cell(self, row, column):
        return _Cell(self._values.get((row, column)), row, column)

    def iter_rows(self, min_row=1, max_row=None, **ignored):
        last = min(max_row or self.ROWS, self.ROWS)
        width = max((c for _, c in self._values), default=0)
        for r in range(min_row, last + 1):
            yield tuple(_Cell(self._values.get((r, c)), r, c) for c in range(1, width + 1))


def _merges_of(archive, sheet_xml):
    """Merged ranges of a sheet, straight from its XML — read-only mode drops them."""
    found = []
    for ref in re.findall(rb'<mergeCell ref="([A-Z]+\d+):([A-Z]+\d+)"', archive.read(sheet_xml)):
        (first, last) = (ref[0].decode(), ref[1].decode())
        first_col, first_row = coordinate_from_string(first)
        last_col, last_row = coordinate_from_string(last)
        found.append(_Range(first_row, last_row, column_index_from_string(first_col),
                            column_index_from_string(last_col)))
    return found


def fast_sheets(source, names):
    """{name: FastSheet} for the day sheets, without parsing the whole workbook."""
    book = load_workbook(source, data_only=True, read_only=True)
    sheets = {}
    try:
        with zipfile.ZipFile(source) as archive:
            for name in names:
                values = {}
                for r, row in enumerate(book[name].iter_rows(min_row=1, max_row=FastSheet.ROWS,
                                                             values_only=True), 1):
                    for c, value in enumerate(row, 1):
                        if value is not None:
                            values[(r, c)] = value
                sheets[name] = FastSheet(values, _merges_of(archive, _sheet_path(archive, name)))
    finally:
        book.close()
    if hasattr(source, "seek"):
        source.seek(0)
    return sheets


def shift_blocks(sheet):
    """[{field: column, "row": header row, "until": where the block ends}] per shift.

    Found by reading the labels, never by fixed coordinates: the workbooks come in
    several versions of the same model — November has the header on row 26 with
    blocks at B/AI/BP, February has row 27 with B/AJ/BR — and a hardcoded layout
    silently drops a whole month.
    """
    spots = sorted((c.row, c.column) for row in sheet.iter_rows(min_row=1, max_row=60)
                   for c in row if isinstance(c.value, str)
                   and c.value.strip().casefold() == SHIFT_FIELD_NAMES["horario"])
    blocks = []
    for i, (row, column) in enumerate(spots):
        until = spots[i + 1][1] if i + 1 < len(spots) and spots[i + 1][0] == row else column + 33
        block = {"row": row, "until": until, "horario": column}
        for field, label in SHIFT_FIELD_NAMES.items():
            for c in range(column, until):
                value = sheet.cell(row, c).value
                if isinstance(value, str) and value.strip().casefold() == label:
                    block[field] = c
                    break
        if "apontado" in block:
            block["diferenca"] = block["apontado"] + 2  # this one has no label of its own
        blocks.append(block)
    return blocks


def is_shift_report(book):
    """Whether this workbook is a shift report with one sheet per day."""
    days = [n for n in book.sheetnames if n.isdigit()]
    return len(days) >= 5 and bool(shift_blocks(book[days[0]]))


def reference_month(book):
    """(year, month, sheet that said so) of the report, or None.

    The day sheets are named "1".."31" and say nothing about the month. Two sheets
    can disagree — the broken "Tabela de Dados" starts at 01/09 while the planning
    sheet lists 01/11..30/11 — so only a sheet that reads like a calendar counts:
    each day exactly once, in a consecutive run.
    """
    for name in (n for n in book.sheetnames if not n.isdigit()):
        days = [v.date() for row in book[name].iter_rows(max_row=60, max_col=30, values_only=True)
                for v in row if isinstance(v, datetime)]
        unique = sorted(set(days))
        if len(unique) >= 5 and len(unique) == len(days) and \
                (unique[-1] - unique[0]).days == len(unique) - 1:
            return unique[0].year, unique[0].month, name
    return None


def _block_field(sheet, block, label):
    """Value of a labelled cell in a shift block's header area, e.g. "SKU DESCRIÇÃO".

    The label sits at a slightly different column in each block (B6, AJ6, BR6), so
    it is searched for instead of hardcoded; the value is the first filled cell to
    its right.
    """
    for r in range(1, block["row"]):
        for c in range(block["horario"], block["until"]):
            cell = sheet.cell(r, c).value
            if isinstance(cell, str) and label in cell.upper():
                for right in range(c + 1, min(c + 8, block["until"])):
                    value = sheet.cell(r, right).value
                    if value not in (None, ""):
                        return str(value).strip()
    return ""


def _block_line(sheet, block):
    """"LINHA D02 - 1º TURNO" over a block -> "D02", or "" when the block has no title."""
    for r in (1, 2, 3):
        for c in range(block["horario"], min(block["horario"] + 4, block["until"])):
            value = sheet.cell(r, c).value
            found = LINE_RE.search(str(value)) if isinstance(value, str) else None
            if found:
                return found.group(1).upper()
    return ""


def _span_label(hours):
    """["06h - 07h", ..., "09h - 10h"] -> "06h - 10h" — one occurrence, four merged hours."""
    first, last = hours[0], hours[-1]
    if first == last:
        return first
    end = last.split("-")[-1].strip() if "-" in last else last
    return f"{first.split('-')[0].strip()} - {end}"


def read_shift_report(book, year, month, sheets=None):
    """(headers, rows) built from the day sheets, with the date rebuilt from the sheet name.

    Sheet "6" is day 6 of the reference month; each sheet holds three shift blocks
    of eight hours side by side. The losses cell is usually MERGED over several
    hour rows — one report covering 06h to 10h — so the row that owns the text
    carries the whole span, and "Horas" lists every hour it covers so the hour
    filter finds it. Days the month does not have are skipped.
    """
    headers = ["Data", "Linha", "Dia", "Turno", "SKU", "Trocas", "Horário", "Período", "Horas",
               "Formato", "Plan", "Produzido", "Apontado", "Diferença", "Principais perdas"]
    rows, number, seen_line = [], 1, ""
    for name in sorted((n for n in book.sheetnames if n.isdigit()), key=int):
        day = int(name)
        try:
            when = date(year, month, day)
        except ValueError:
            continue  # 31 in a 30-day month
        sheet = sheets[name] if sheets else book[name]
        blocks = shift_blocks(sheet)
        # only the first shift is usually titled "LINHA D01 - 1º TURNO"; the others
        # belong to the same line, and leaving them nameless threw them into a
        # separate block at the end of the list, as if the day stopped at 14h
        titles = [_block_line(sheet, block) for block in blocks]
        seen_line = next((t for t in titles if t), seen_line)
        for shift, block in enumerate(blocks, 1):
            first_row = block["row"] + 1
            last_row = block["row"] + SHIFT_HOURS
            line = titles[shift - 1] or seen_line or "—"
            sku = _block_field(sheet, block, "SKU")
            swaps = [t for i in (1, 2, 3)
                     if (t := _block_field(sheet, block, f"TROCA {i}"))
                     and t.upper() not in ("N/A", "NA", "-", "0")]
            losses = block.get("perdas")
            merges = {r.min_row: min(r.max_row, last_row) for r in sheet.merged_cells.ranges
                      if losses and r.min_col <= losses <= r.max_col and r.min_row >= first_row}
            hour_of = {r: str(sheet.cell(r, block["horario"]).value or "").strip()
                       for r in range(first_row, last_row + 1)}
            for r in range(first_row, last_row + 1):
                cell = lambda field: (sheet.cell(r, block[field]).value
                                      if field in block else None)
                if not hour_of[r]:
                    continue
                covered = [hour_of[x] for x in range(r, merges.get(r, r) + 1) if hour_of.get(x)]
                number += 1
                rows.append((number, [when, line, day, shift, sku, " · ".join(swaps),
                                      hour_of[r], _span_label(covered),
                                      " | ".join(covered), cell("formato"), cell("plan"),
                                      cell("produzido"), cell("apontado"), cell("diferenca"),
                                      cell("perdas")]))
    return headers, rows


MONTH_NAMES = ("janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
               "agosto", "setembro", "outubro", "novembro", "dezembro")


def month_from_name(name):
    """(year or None, month) read from a file name, or None.

    The workbooks in the bundle are named by month — "Relatório de Turno - Linha 1-
    Março.xlsx" — and carry no year, so the year has to come from somewhere else.
    """
    plain = unicodedata.normalize("NFKD", str(name).lower()).encode("ascii", "ignore").decode()
    stamped = re.search(r"(20\d{2})[-_/.]?(\d{1,2})|(\d{1,2})[-_/.](20\d{2})", plain)
    if stamped:
        year, month = (stamped.group(1), stamped.group(2)) if stamped.group(1) \
            else (stamped.group(4), stamped.group(3))
        if 1 <= int(month) <= 12:
            return int(year), int(month)
    for number, month_name in enumerate(MONTH_NAMES, 1):
        plain_name = unicodedata.normalize("NFKD", month_name).encode("ascii", "ignore").decode()
        if re.search(rf"\b{plain_name}\b", plain):
            return None, number
    return None


PLACEHOLDER_YEAR = 2000  # leap, so 29/02 survives until the real year is known


def _with_source(headers, rows, name):
    """Tag every row with the file it came from, so a day can say which report it is."""
    return headers + ["Arquivo"], [(number, list(values) + [name]) for number, values in rows]


def _with_year(headers, rows, year):
    """Rewrite the date of every row to the year that was only known at the end."""
    date_at = _at(headers, "Data")
    if date_at is None:
        return rows
    fixed = []
    for number, values in rows:
        day = values[date_at]
        try:
            values = list(values)
            values[date_at] = date(year, day.month, day.day)
        except ValueError:
            continue  # 29/02 in a year that has no 29/02
        fixed.append((number, values))
    return fixed


def load_bundle(path, progress=None):
    """Every .xlsx inside a zip read as one dataset, each with its own month.

    Each workbook names its own month (planning sheet, or the month in the file
    name), so the screen has no month field to fill: the date range does that job.
    """
    parts, months, headers = [], [], []
    with zipfile.ZipFile(path) as bundle:
        inside = sorted(n for n in bundle.namelist()
                        if n.lower().endswith(".xlsx") and not n.startswith("__MACOSX")
                        and not Path(n).name.startswith("~$"))
        if not inside:
            raise ValueError("O zip não tem nenhuma planilha .xlsx dentro.")
        for i, name in enumerate(inside, 1):
            short = Path(name).name
            if progress:
                progress(i, len(inside), short)
            data = io.BytesIO(bundle.read(name))
            book = load_workbook(data, data_only=True, read_only=True)
            try:
                if not is_shift_report(book):
                    months.append((short, None, "não é relatório de turno"))
                    continue
                found = reference_month(book)
                named = month_from_name(short)
                if found and named and named[1] != found[1]:
                    # the file name and the planning sheet disagree: the name is chosen
                    # per report, the sheet is often left over from the template. The
                    # name wins and the conflict is said out loud, never swallowed.
                    when = (named[0], named[1])
                    source = "nome do arquivo (a aba {} dizia {:02d}/{})".format(
                        found[2], found[1], found[0])
                elif found:
                    when, source = found[:2], f"aba {found[2]}"
                elif named and named[0]:
                    when, source = named, "nome do arquivo"
                elif named:
                    when, source = (None, named[1]), "nome do arquivo (sem ano)"
                else:
                    when, source = (None, None), "não encontrado"
                months.append((short, when, source))
                # the year can still be missing here — it is borrowed from the other
                # workbooks after the loop — so read with a placeholder and remap below
                days = [n for n in book.sheetnames if n.isdigit()]
                headers, rows = read_shift_report(book, when[0] or PLACEHOLDER_YEAR, when[1],
                                                  sheets=fast_sheets(data, days))
                headers, rows = _with_source(headers, rows, short)
                parts.append((short, headers, rows))
            finally:
                book.close()
    # a workbook that only named the month borrows the year the others agree on
    years = Counter(when[0] for _, when, _ in months if when and when[0])
    year = years.most_common(1)[0][0] if years else date.today().year
    borrowed = {short for short, when, _ in months if when and when[1] and not when[0]}
    for i, (short, when, source) in enumerate(months):
        if short in borrowed:
            months[i] = (short, (year, when[1]), source + f", ano {year} dos outros arquivos")
    if borrowed:
        parts = [(short, part_headers, _with_year(part_headers, rows, year)
                  if short in borrowed else rows)
                 for short, part_headers, rows in parts]
    return parts, months, headers


HISTORY_FILE = Path.home() / ".resumo-ocorrencias" / "historico.json"
HISTORY_LIMIT = 15


def history():
    """Spreadsheets opened before, newest first. A missing or broken file is no history."""
    try:
        saved = json.loads(HISTORY_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    return [e for e in saved if isinstance(e, dict) and e.get("path")]


def remember(path, entry):
    """Record this spreadsheet at the top of the history and keep the file small."""
    kept = [e for e in history() if e["path"] != str(path)]
    kept.insert(0, {"path": str(path), "when": datetime.now().isoformat(timespec="seconds"), **entry})
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    HISTORY_FILE.write_text(json.dumps(kept[:HISTORY_LIMIT], ensure_ascii=False, indent=1),
                            encoding="utf-8")
    return kept[:HISTORY_LIMIT]


def forget(path):
    """Take one spreadsheet out of the history and give back what is left."""
    kept = [e for e in history() if e["path"] != str(path)]
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    HISTORY_FILE.write_text(json.dumps(kept, ensure_ascii=False, indent=1), encoding="utf-8")
    return kept


def forget_all():
    """Wipe the history."""
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    HISTORY_FILE.write_text("[]", encoding="utf-8")
    return []


# --- Admin access ----------------------------------------------------------------
ADMIN_ROUNDS = 200_000


def admin_file():
    """admin.json: bundled into the .exe at build time, next to the script from source.

    Never committed — the repository is public, and even a salted hash of a
    short password is crackable once it is on display.
    """
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / "admin.json"


def make_admin(user, password, path=None):
    """Write the admin's login with a salted PBKDF2 hash; the password is never stored."""
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, ADMIN_ROUNDS)
    target = Path(path or admin_file())
    target.write_text(json.dumps({"user": user.strip(), "salt": salt.hex(), "hash": digest.hex(),
                                  "rounds": ADMIN_ROUNDS}), encoding="utf-8")
    return target


def is_admin(user, password, path=None):
    """Whether this login and password are the admin's. No admin file means no admin.

    ponytail: a check inside a desktop .exe keeps the admin buttons away from the
    floor; it does not stop someone who unpacks the .exe. Real protection is a
    server deciding who may write the base.
    """
    try:
        saved = json.loads(Path(path or admin_file()).read_text(encoding="utf-8"))
        salt, expected = bytes.fromhex(saved["salt"]), bytes.fromhex(saved["hash"])
        rounds, known = int(saved["rounds"]), str(saved["user"])
    except (OSError, ValueError, KeyError, TypeError):
        return False
    typed = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, rounds)
    return hmac.compare_digest(typed, expected) and user.strip().casefold() == known.casefold()


# --- Saved base: what the admin fed, kept until the admin deletes it --------------
BASE_DIR = HISTORY_FILE.parent / "base"
STORE_LOCK = threading.RLock()  # the polling thread and the screen both write here
ENTRY_ID = re.compile(r"[0-9a-f]{16}")


def _encode(value):
    """A cell as JSON: dates travel tagged and come back as dates."""
    if isinstance(value, datetime):
        return {"$dt": value.isoformat()}
    if isinstance(value, date):
        return {"$d": value.isoformat()}
    return str(value)  # a time or anything exotic keeps its text, never turns into nothing


def _decode(obj):
    if "$dt" in obj:
        return datetime.fromisoformat(obj["$dt"])
    if "$d" in obj:
        return date.fromisoformat(obj["$d"])
    return obj


def _write_json(path, payload):
    """Write through a temp file, so a crash mid-write never leaves half a file."""
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(json.dumps(payload, default=_encode, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def base_entries(folder=None):
    """One entry per saved workbook. A missing or broken index is an empty base.

    Ids are checked before anything uses them as a file name: an edited index
    must not be able to point outside the folder.
    """
    try:
        saved = json.loads((Path(folder or BASE_DIR) / "indice.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return []
    if not isinstance(saved, list):
        return []
    return [e for e in saved if isinstance(e, dict) and e.get("key")
            and ENTRY_ID.fullmatch(str(e.get("id", "")))]


def _workbook_key(headers, rows, name):
    """(key, line, (year, month)) of a workbook, read from its rows: one line in one month.

    The file name is not the identity — a browser saves the same zip again as
    "(1)", and the same month sent again has to replace the old one, not double
    every total. Only a workbook whose line was never written falls back to it.
    """
    line_at, date_at = _at(headers, "Linha"), _at(headers, "Data")
    lines = Counter(str(v[line_at]) for _, v in rows
                    if line_at is not None and v[line_at] not in (None, "", "—"))
    months = Counter((d.year, d.month) for d in
                     (parse_date(v[date_at]) for _, v in rows if date_at is not None) if d)
    line = lines.most_common(1)[0][0] if lines else ""
    month = months.most_common(1)[0][0] if months else None
    when = f"{month[0]}-{month[1]:02d}" if month else "?"
    return f"{line or name}|{when}", line, month


def base_add(data, origin, folder=None):
    """Save every workbook of a loaded dataset into the base. (saved, replaced) file names.

    The same line and month replaces what was there. Only shift reports with a
    month that was read go in: the base is one table, and a guessed month would
    file a whole month of production under the wrong date for good.
    """
    if data.get("kind") != "turno":
        raise ValueError("Só relatório de turno entra na base.")
    if data.get("month_guessed"):
        raise ValueError("Nenhuma aba diz o mês desta planilha — envie dentro do zip, com o mês"
                         " no nome do arquivo.")
    folder = Path(folder or BASE_DIR)
    folder.mkdir(parents=True, exist_ok=True)
    headers, file_at = data["headers"], _at(data["headers"], "Arquivo")
    parts = {}
    for number, values in data["rows"]:
        name = values[file_at] if file_at is not None else origin
        parts.setdefault(str(name), []).append((number, values))
    with STORE_LOCK:
        entries, stale, saved, replaced = base_entries(folder), [], [], []
        for name, rows in parts.items():
            key, line, month = _workbook_key(headers, rows, name)
            for old in [e for e in entries if e["key"] == key]:
                entries.remove(old)
                stale.append(old["id"])
                replaced.append(old.get("file", ""))
            entry = {"id": secrets.token_hex(8), "key": key, "file": name, "line": line,
                     "month": list(month) if month else None, "rows": len(rows),
                     "origin": origin, "added": datetime.now().isoformat(timespec="seconds")}
            _write_json(folder / f"{entry['id']}.json",
                        {"headers": headers, "rows": [values for _, values in rows]})
            entries.append(entry)
            saved.append(name)
        _write_json(folder / "indice.json", entries)  # new files exist before the index says so
        for gone in stale:  # and the old ones go only after the index stopped pointing at them
            (folder / f"{gone}.json").unlink(missing_ok=True)
    return saved, replaced


def base_remove(ids, folder=None):
    """Take workbooks out of the base for good; give back what is left."""
    folder = Path(folder or BASE_DIR)
    gone = {i for i in ids if ENTRY_ID.fullmatch(str(i))}
    with STORE_LOCK:
        kept = [e for e in base_entries(folder) if e["id"] not in gone]
        if folder.exists():
            _write_json(folder / "indice.json", kept)
        for i in gone:
            (folder / f"{i}.json").unlink(missing_ok=True)
    return kept


def base_load(folder=None):
    """The saved base as one dataset for the screen, or None when it is empty.

    A workbook whose file is gone or unreadable is listed in "missing", so the
    screen can say it instead of showing smaller totals as if they were whole.
    """
    folder = Path(folder or BASE_DIR)
    entries = sorted(base_entries(folder),
                     key=lambda e: (str(e.get("line") or ""), e.get("month") or [0, 0]))
    headers, rows, months, missing = None, [], [], []
    for entry in entries:
        try:
            saved = json.loads((folder / f"{entry['id']}.json").read_text(encoding="utf-8"),
                               object_hook=_decode)
            part_headers, part_rows = saved["headers"], saved["rows"]
        except (OSError, ValueError, KeyError, TypeError):
            missing.append(entry.get("file", entry["id"]))
            continue
        if headers is None:
            headers = part_headers
        elif part_headers != headers:
            missing.append(entry.get("file", entry["id"]))
            continue
        rows.extend((len(rows) + i, values) for i, values in enumerate(part_rows, 1))
        month = tuple(entry["month"]) if entry.get("month") else None
        months.append((entry.get("file", ""), month, "base salva"))
    if not rows:
        return {"missing": missing} if missing else None
    read_months = [when for _, when, _ in months if when]
    return {"kind": "turno", "sheet": "base salva", "headers": headers, "rows": rows,
            "path": "", "month": min(read_months) if read_months else None, "month_from": None,
            "month_guessed": False, "months": months, "base": True, "missing": missing,
            "columns": dict(SHIFT_COLUMNS)}


# --- Feeding from SharePoint: Graph, a plain link, or a folder OneDrive keeps in sync --
LOGIN_WALL = ("O link não devolveu um .zip nem um .xlsx — o SharePoint pediu login. Cole o token"
              " do Graph em Adicionar API, use um link que baixe sem login, ou Enviar .zip.")
GRAPH = "https://graph.microsoft.com/v1.0"
GRAPH_HOST = "https://graph.microsoft.com/"
SPREADSHEETS = (".xlsx", ".zip")
WALK_LIMIT = 400  # files: a key at the root of a whole library would pull all of it
GRAPH_PATHS = ("drives/", "sites/", "me/drive/")
GRAPH_ERRORS = {401: "O token expirou ou não vale — cole um novo em Adicionar API.",
                403: "O token não dá acesso a esse arquivo.",
                404: "Nenhum arquivo com essa chave no SharePoint."}
SOURCES_FILE = "fontes.json"
SYNC_EVERY_MS = 60_000  # a minute: fresh enough for a shift report, one light call per file


def direct_link(url):
    """A SharePoint share link answers with a page; download=1 asks for the file itself."""
    parts = urllib.parse.urlsplit(url.strip())
    if parts.scheme not in ("http", "https") or not parts.hostname:
        raise ValueError("O link precisa começar com https://")
    if parts.hostname.endswith(".sharepoint.com"):
        query = urllib.parse.parse_qsl(parts.query, keep_blank_values=True)
        if not any(name == "download" for name, _ in query):
            query.append(("download", "1"))
        parts = parts._replace(query=urllib.parse.urlencode(query))
    return urllib.parse.urlunsplit(parts)


def share_id(link):
    """Graph's id for a sharing link: "u!" + the link in unpadded base64url (Graph docs)."""
    return "u!" + base64.urlsafe_b64encode(link.strip().encode()).decode().rstrip("=")


def graph_key(key):
    """The drives/... path inside whatever was pasted: a bare key, or a full Graph URL
    ending in /content or /children — the forms Graph Explorer and Copilot hand out."""
    key = key.strip()
    if key[:4].upper() == "GET ":
        key = key[4:].strip()
    if key.startswith(GRAPH_HOST):
        key = key[len(GRAPH_HOST):].split("?", 1)[0]
        key = key.split("/", 1)[1] if key.startswith(("v1.0/", "beta/")) else key
    for tail in ("/content", "/children"):
        if key.endswith(tail):
            key = key[:-len(tail)]
    return key.lstrip("/")


def graph_item_url(key):
    """Where Graph describes what a key points to: a drives/... path or a SharePoint link."""
    path = graph_key(key)
    if path.startswith(GRAPH_PATHS):
        return f"{GRAPH}/{path}"
    if key.strip().startswith(("http://", "https://")):
        return f"{GRAPH}/shares/{share_id(key)}/driveItem"
    raise ValueError("Com o token, a chave é drives/<id>/items/<id> (ou a URL do Graph com ela),"
                     " ou o link do arquivo no SharePoint.")


def month_of(names):
    """(year, month) that folder and file names say — "2026" / "10 - Outubro" — or None.

    The deepest name naming a month wins; its year is the one in that same name
    or the nearest one above it. Without both, the workbook's own sheets decide.
    """
    for i in range(len(names) - 1, -1, -1):
        said = month_from_name(names[i])
        if not said:
            continue
        year = said[0]
        for above in reversed(names[:i + 1]):
            if year:
                break
            found = re.search(r"\b(20\d{2})\b", str(above))
            year = int(found.group(1)) if found else None
        return (year, said[1]) if year else None
    return None


def _graph_json(url, token, timeout=30):
    """The driveItem Graph describes, as a dict; every failure said in words."""
    request = urllib.request.Request(url, headers={"Authorization": f"Bearer {token.strip()}",
                                                   "Accept": "application/json"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as answer:
            payload = json.loads(answer.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        raise ValueError(GRAPH_ERRORS.get(error.code, f"O Graph respondeu {error.code}.")) from error
    except urllib.error.URLError as error:
        raise ValueError(f"Não deu para falar com o Graph: {error.reason}") from error
    except ValueError as error:
        raise ValueError("O Graph respondeu algo que não é JSON.") from error
    if not isinstance(payload, dict):  # "[]" and "7" parse as JSON too
        raise ValueError("O Graph respondeu algo que não é a descrição de um arquivo.")
    return payload


def _open(request, timeout):
    """urlopen with every failure said in words; None when the server answers 304."""
    try:
        return urllib.request.urlopen(request, timeout=timeout)
    except urllib.error.HTTPError as error:
        if error.code == 304:
            return None
        raise ValueError(LOGIN_WALL if error.code in (401, 403)
                         else f"O link respondeu {error.code}.") from error
    except urllib.error.URLError as error:
        raise ValueError(f"Não deu para acessar o link: {error.reason}") from error


def _save(request, folder, name=None, timeout=60):
    """Stream a download into `folder`: (local .zip/.xlsx, {"etag", "modified"}), or
    (None, {}) when the server answers 304 — nothing new since the version sent.

    Streamed — a month of reports is hundreds of MB. A link that is not public
    answers with a login page instead of the file: that is said, never handed to
    the reader as if it were a spreadsheet.
    """
    partial = Path(folder) / "download.part"
    answer = _open(request, timeout)
    if answer is None:
        return None, {}
    with answer, open(partial, "wb") as out:
        head = answer.read(4)
        if head != b"PK\x03\x04":  # every .zip and .xlsx starts like this; a page does not
            raise ValueError(LOGIN_WALL)
        out.write(head)
        while chunk := answer.read(1 << 20):
            out.write(chunk)
        tags = {"etag": answer.headers.get("ETag", ""),
                "modified": answer.headers.get("Last-Modified", "")}
        name = name or answer.headers.get_filename() or urllib.parse.urlsplit(answer.geturl()).path
    try:
        with zipfile.ZipFile(partial) as package:
            workbook = "xl/workbook.xml" in package.namelist()
    except zipfile.BadZipFile as error:
        raise ValueError("O arquivo baixado veio quebrado.") from error
    stem = re.sub(r'[<>:"/\\|?*]', "_", Path(urllib.parse.unquote(name or "")).stem) or "sharepoint"
    return partial.replace(Path(folder) / (stem + (".xlsx" if workbook else ".zip"))), tags


def _who(item):
    """Who last changed a Graph item (lastModifiedBy.user.displayName), or ""."""
    changed_by = item.get("lastModifiedBy")
    user = changed_by.get("user") if isinstance(changed_by, dict) else None
    return str(user.get("displayName") or "") if isinstance(user, dict) else ""


def _graph_names(item):
    """The folder names above a Graph item (parentReference.path, percent-encoded) + its own."""
    path = str((item.get("parentReference") or {}).get("path") or "")
    above = path.split("root:", 1)[1] if "root:" in path else ""
    return [urllib.parse.unquote(part) for part in above.split("/") if part] + [str(item.get("name") or "")]


def _graph_item_of(item):
    reference = item.get("parentReference") or {}
    if not reference.get("driveId") or not item.get("id"):
        raise ValueError("O Graph não disse onde fica esse item (driveId).")
    return f"{GRAPH}/drives/{reference['driveId']}/items/{item['id']}"


def _graph_files(folder_item, token, names, timeout, depth=0, found=None):
    """[(file item, names)] under a Graph folder — ANO → MÊS → ARQUIVO, four levels down.

    Pages are followed through @odata.nextLink (200 items a page), but only on
    the Graph host: the token never travels anywhere else.
    """
    found = [] if found is None else found
    url = _graph_item_of(folder_item) + "/children"
    while url:
        if not url.startswith(GRAPH):
            raise ValueError("O Graph mandou continuar a lista fora do Graph — parei.")
        page = _graph_json(url, token, timeout)
        children = page.get("value") if isinstance(page.get("value"), list) else []
        for child in (c for c in children if isinstance(c, dict)):
            name = str(child.get("name") or "")
            if "folder" in child:
                if depth < 3:
                    _graph_files(child, token, names + [name], timeout, depth + 1, found)
            elif name.lower().endswith(SPREADSHEETS) and not name.startswith("~$"):
                found.append((child, names + [name]))
        if len(found) > WALK_LIMIT:
            raise ValueError(f"Mais de {WALK_LIMIT} planilhas embaixo dessa chave — use a pasta"
                             " do ano ou do mês.")
        url = page.get("@odata.nextLink")
    return found


def source_versions(source):
    """{file id: version} a source was last seen at (builds before folders kept just one)."""
    versions = source.get("versions")
    if isinstance(versions, dict):
        return {str(k): str(v) for k, v in versions.items()}
    return {"": str(source["version"])} if source.get("version") else {}


def check_source(source, token="", folder=None, timeout=60):
    """What one live source holds now: [(file id, version, fetch or None, who, month, name)].

    A file key gives one entry; a folder key — the year, or a month — is walked
    down to its spreadsheets, ANO → MÊS → ARQUIVO, the month read from the folder
    names ("2026" / "10 - Outubro") instead of guessed. `fetch` is there only for
    files whose version moved — the content tag Graph gives (cTag), the ETag a
    plain link answers with, or the time and size of a synced file — and it
    downloads when called, so one bad file never stops its siblings.
    """
    key = str(source.get("key") or "").strip()
    if not key:
        raise ValueError("Fonte sem chave.")
    known, folder = source_versions(source), Path(folder or tempfile.mkdtemp())
    graph = graph_key(key).startswith(GRAPH_PATHS)
    if not key.startswith(("http://", "https://")) and not graph:
        return _check_local(Path(key), known)
    if token.strip():
        return _check_graph(key, token, known, folder, timeout)
    if graph:
        raise ValueError("Essa chave é do Graph (drives/...): cole o token em Adicionar API.")
    return [_check_link(key, known.get("", ""), folder, timeout)]


def _check_local(path, known):
    """A file, or a folder, the OneDrive client keeps in sync: versions from time and size."""
    if path.is_dir():
        files = sorted(p for p in path.rglob("*") if p.is_file()
                       and p.suffix.lower() in SPREADSHEETS and not p.name.startswith("~$"))
        if len(files) > WALK_LIMIT:
            raise ValueError(f"Mais de {WALK_LIMIT} planilhas nessa pasta — use a do ano ou a do mês.")
    elif path.is_file():
        files = [path]
    else:
        raise ValueError(f"Arquivo ou pasta não encontrado: {path}")
    found = []
    for file in files:
        stat = file.stat()
        version = f"arquivo:{stat.st_mtime_ns}:{stat.st_size}"
        inside = file != path
        fid = file.relative_to(path).as_posix() if inside else ""
        # only the names from the key down (plus its parent, for the year): a folder
        # of the user's — "Maio Silva" — must not pass for a month
        names = ([path.parent.name, path.name, *file.relative_to(path).parts] if inside
                 else list(file.parts[-4:]))
        found.append((fid, version, (lambda f=file: f) if known.get(fid) != version else None,
                      "", month_of(names), file.name))
    return found


def _check_graph(key, token, known, folder, timeout):
    """Through Microsoft Graph: the token goes to Graph only, never to a download address."""
    item = _graph_json(graph_item_url(key), token, timeout)
    names = _graph_names(item)
    if "folder" in item:
        files = [(str(child.get("id") or child_names[-1]), child, child_names)
                 for child, child_names in _graph_files(item, token, names, timeout)]
    else:
        files = [("", item, names)]

    def downloader(child, target):
        def fetch():
            url = child.get("@microsoft.graph.downloadUrl")
            if not url:  # a listing may leave it out: ask the item itself
                url = _graph_json(_graph_item_of(child), token, timeout).get(
                    "@microsoft.graph.downloadUrl")
            if not url:
                raise ValueError("O Graph não deu o endereço de download — é uma pasta?")
            target.mkdir(parents=True, exist_ok=True)
            # the download address is pre-signed: the token must not travel with it
            return _save(urllib.request.Request(str(url)), target, child.get("name"), timeout)[0]
        return fetch

    found = []
    for index, (fid, child, child_names) in enumerate(files):
        version = "graph:" + str(child.get("cTag") or child.get("eTag")
                                 or child.get("lastModifiedDateTime") or "")
        fetch = downloader(child, folder / str(index)) if known.get(fid) != version else None
        found.append((fid, version, fetch, _who(child), month_of(child_names),
                      str(child.get("name") or "")))
    return found


def _check_link(key, known, folder, timeout):
    """A plain link: asks with the version it has; the server's 304 means nothing new."""
    headers = {"User-Agent": "Mozilla/5.0"}
    if known.startswith("etag:"):
        headers["If-None-Match"] = known[len("etag:"):]
    elif known.startswith("modificado:"):
        headers["If-Modified-Since"] = known[len("modificado:"):]
    path, tags = _save(urllib.request.Request(direct_link(key), headers=headers), folder,
                       timeout=timeout)
    if path is None:  # 304: the same version as last time
        return "", known, None, "", None, key
    if tags["etag"]:
        version = "etag:" + tags["etag"]
    elif tags["modified"]:
        version = "modificado:" + tags["modified"]
    else:  # a server that says nothing about versions: the content decides
        with open(path, "rb") as content:
            version = "sha256:" + hashlib.file_digest(content, "sha256").hexdigest()
    if version == known:
        path.unlink()
        return "", known, None, "", None, path.name
    return "", version, (lambda p=path: p), "", month_of([path.name]), path.name


def _month_text(month):
    return f"{month[1]:02d}/{month[0]}" if month else "mês pelo conteúdo da planilha"


def probe_source(key, token="", timeout=30):
    """What a key reaches, read only: [lines]. Nothing is downloaded or saved.

    "Testar conexão" answers with this: whether the token and the key open the
    file or the folder, and which spreadsheets and months are behind it — and
    the base, the sources and SharePoint stay exactly as they were.
    """
    key = key.strip()
    if not key:
        raise ValueError("Cole a chave para testar.")
    graph = graph_key(key).startswith(GRAPH_PATHS)
    if not key.startswith(("http://", "https://")) and not graph:
        entries = _check_local(Path(key), {})
        return [f"Achei {len(entries)} planilha(s):"] + [
            f"{name} · {_month_text(month)}" for _, _, _, _, month, name in entries]
    if token.strip():
        item = _graph_json(graph_item_url(key), token, timeout)
        names = _graph_names(item)
        if "folder" not in item:
            return ["Conectou no arquivo {} · {} · modificado por {} em {}".format(
                item.get("name"), _month_text(month_of(names)), _who(item) or "?",
                str(item.get("lastModifiedDateTime") or "?")[:16].replace("T", " "))]
        files = _graph_files(item, token, names, timeout)
        return [f"Conectou na pasta {item.get('name')}: {len(files)} planilha(s)"] + [
            "{} · {} · modificado por {}".format(child_names[-1], _month_text(month_of(child_names)),
                                                 _who(child) or "?")
            for child, child_names in files]
    if graph:
        raise ValueError("Essa chave é do Graph (drives/...): cole o token para testar.")
    answer = _open(urllib.request.Request(direct_link(key), headers={"User-Agent": "Mozilla/5.0"}),
                   timeout)
    with answer:  # only the first bytes: enough to know it is a file, not a login page
        if answer.read(4) != b"PK\x03\x04":
            raise ValueError(LOGIN_WALL)
        name = answer.headers.get_filename() or urllib.parse.urlsplit(answer.geturl()).path
    return [f"O link entrega um .zip/.xlsx: {Path(urllib.parse.unquote(name)).name}"]


def load_sources(folder=None):
    """{"token", "sources"} that Adicionar API keeps. A missing or broken file is empty."""
    try:
        saved = json.loads((Path(folder or BASE_DIR) / SOURCES_FILE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        saved = {}
    saved = saved if isinstance(saved, dict) else {}
    sources = saved.get("sources") if isinstance(saved.get("sources"), list) else []
    return {"token": str(saved.get("token") or ""),
            "sources": [s for s in sources if isinstance(s, dict) and s.get("key")
                        and ENTRY_ID.fullmatch(str(s.get("id", "")))]}


def _change_sources(change, folder=None):
    """Read, change and write fontes.json under the store lock; the new state."""
    folder = Path(folder or BASE_DIR)
    with STORE_LOCK:
        state = load_sources(folder)
        change(state)
        folder.mkdir(parents=True, exist_ok=True)
        _write_json(folder / SOURCES_FILE, state)
    return state


def add_source(label, key, folder=None):
    """Follow a file from now on; it is read on the next check."""
    key = key.strip()
    if not key:
        raise ValueError("Cole a chave: o link do arquivo, um caminho drives/..., ou o arquivo"
                         " da pasta que o OneDrive sincroniza.")
    if key.startswith(("http://", "https://")):
        direct_link(key)  # says what is wrong with the link before it is saved
    entry = {"id": secrets.token_hex(8), "label": label.strip() or "fonte", "key": key,
             "version": "", "checked": "", "changed": "", "by": "", "problem": ""}
    return _change_sources(lambda state: state["sources"].append(entry), folder)


def remove_source(source_id, folder=None):
    """Stop following a file. What it already brought stays in the base."""
    return _change_sources(lambda state: state.update(
        sources=[s for s in state["sources"] if s["id"] != source_id]), folder)


def set_token(token, folder=None):
    return _change_sources(lambda state: state.update(token=token.strip()), folder)


def sync_sources(folder=None):
    """Check every live source once and feed the base with what changed: [(label, who)].

    One broken source, or one broken file inside a folder, never stops the rest;
    the problem stays on the source for the screen. A file's version only moves
    after it was read, so one that could not be read is fetched again next round;
    one the base refuses (not a shift report) waits until it changes.
    """
    state, results, updated = load_sources(folder), {}, []
    for source in state["sources"]:
        label, temp = source.get("label") or "fonte", tempfile.mkdtemp()
        now = datetime.now().isoformat(timespec="seconds")
        try:
            entries = check_source(source, state["token"], temp)
        except Exception as error:  # the whole source failed: it keeps what it knew
            results[source["id"]] = {"checked": now, "problem": str(error) or type(error).__name__}
            shutil.rmtree(temp, ignore_errors=True)
            continue
        known, ignored = source_versions(source), dict(source.get("ignored") or {})
        versions, problems, who_changed = {}, [], None
        for fid, version, fetch, who, month, name in entries:
            if fetch is None:
                versions[fid] = version
                continue
            try:
                data = load(fetch(), month)
            except Exception as error:  # unreadable now — mid-upload, maybe: next round again
                problems.append(f"{name}: {error}")
                if fid in known:
                    versions[fid] = known[fid]
                continue
            try:
                base_add(data, f"{label} · {name}", folder)
                ignored.pop(fid, None)
                who_changed = who if who_changed is None or who else who_changed
            except ValueError as error:  # read, but not for the base: wait for it to change
                ignored[fid] = f"{name}: {error}"
            versions[fid] = version
        result = {"checked": now, "versions": versions, "files": len(entries),
                  "ignored": {k: v for k, v in ignored.items() if k in versions},
                  "problem": "; ".join(problems)}
        if who_changed is not None:
            result.update(changed=now, by=who_changed)
            updated.append((label, who_changed))
        shutil.rmtree(temp, ignore_errors=True)
        results[source["id"]] = result

    def merge(fresh):  # into the freshest file: a source added during the check survives
        for source in fresh["sources"]:
            source.update(results.get(source["id"], {}))

    _change_sources(merge, folder)
    return updated


def load(path, month=None, progress=None):
    """Open the file once and give back everything the screen needs.

    A .zip is every shift report inside it, merged. A single shift report is read
    from the day sheets (its "Tabela de Dados" has broken dates); anything else
    falls back to the plain table reader.
    """
    if str(path).lower().endswith(".zip"):
        parts, months, headers = load_bundle(path, progress)
        rows, number = [], 0
        for short, part_headers, part_rows in parts:
            for _, values in part_rows:
                number += 1
                rows.append((number, values))
        read_months = [when for _, when, _ in months if when and all(when)]
        return {"kind": "turno", "sheet": f"{len(parts)} planilhas do zip", "headers": headers,
                "rows": rows, "month": min(read_months) if read_months else None,
                "month_from": None, "month_guessed": False, "months": months, "path": str(path),
                "columns": dict(SHIFT_COLUMNS)}
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        if is_shift_report(book):
            found = reference_month(book)
            when = month or (found[:2] if found else (date.today().year, date.today().month))
            days = [n for n in book.sheetnames if n.isdigit()]
            headers, rows = read_shift_report(book, *when, sheets=fast_sheets(path, days))
            headers, rows = _with_source(headers, rows, Path(path).name)
            return {"kind": "turno", "sheet": "abas por dia", "headers": headers, "rows": rows,
                    "path": str(path), "month": when,
                    "month_from": found[2] if found and not month else None,
                    # no sheet named the month and nobody asked for one: it is a guess,
                    # and a guessed date must say so instead of passing for a read value
                    "month_guessed": not month and not found,
                    "columns": dict(SHIFT_COLUMNS)}
        book.close()  # the plain table reader needs random access, so it opens its own
        book = load_workbook(path, data_only=True)
        sheet, headers, rows = read(path, book=book)
        date_col, hour_col, text_col = detect_columns(headers, rows)
        return {"kind": "tabela", "sheet": sheet, "headers": headers, "rows": rows,
                "path": str(path), "month": None, "month_from": None, "month_guessed": False,
                "columns": {"date": date_col, "hour": hour_col, "label": hour_col,
                            "text": text_col, "line": "Linha" if "Linha" in headers else ""}}
    finally:
        book.close()


def sentences_with(text, terms):
    """Only the sentences carrying one of the terms — a cell holds a whole shift log."""
    if not terms:
        return str(text).strip()
    parts = re.split(r"(?<=[.;!?])\s+|\n+", str(text))
    hits = [p.strip() for p in parts if p.strip() and any(t in p.casefold() for t in terms)]
    return "\n".join(hits) if hits else str(text).strip()


def distinct(headers, rows, column):
    """Sorted distinct values of a column — feeds the hour picker."""
    i = _column_index(headers, column) - 1
    return sorted({str(v[i]).strip() for _, v in rows if v[i] is not None and str(v[i]).strip()})


def detect_columns(headers, rows):
    """Find the date, hour and occurrence columns: (date_col, hour_col, text_col).

    Goes by content, not by header name — otherwise "Hora" (1..24) would beat
    "Horário" ("07h - 08h") in this spreadsheet.
    """
    def score(i, test):
        return sum(1 for _, v in rows if v[i] is not None and str(v[i]).strip() and test(v[i]))

    def best(scores):  # needs half the rows, so it never guesses a random column
        i, points = max(scores.items(), key=lambda x: x[1], default=(None, 0))
        return headers[i] if points >= max(1, len(rows) // 2) else ""

    used = [i for i in range(len(headers)) if headers[i]]
    dates = {i: score(i, lambda v: parse_date(v) is not None) for i in used}
    hours = {i: score(i, lambda v: bool(HOUR_RE.fullmatch(str(v)))) for i in used}
    # ponytail: the occurrence is the longest text column. If two long text
    # columns ever show up, this becomes a picker on screen.
    length = {i: sum(len(str(v[i])) for _, v in rows if isinstance(v[i], str)) for i in used}
    text = max(length.items(), key=lambda x: x[1], default=(None, 0))
    return best(dates), best(hours), headers[text[0]] if text[1] > 40 else ""


def _column_index(headers, name):
    found = [i for i, h in enumerate(headers, 1) if h == name]
    if len(found) != 1:
        raise ValueError(f"Coluna {name!r}: {len(found)} colunas com esse nome no cabeçalho.")
    return found[0]


def _matcher(headers, text_col, terms, date_col, start, end, hour_col, shift,
             line_col=None, lines=()):
    """One sieve, every view filters through it — never two rulesets.

    `shift` is 1, 2 or 3 (see SHIFTS), read from the hour column of each row.
    """
    if start and end and start > end:
        raise ValueError("Data inicial maior que a data final.")
    if terms and not text_col:
        raise ValueError("A planilha não tem uma coluna de texto para procurar a palavra.")
    if (start or end) and not date_col:
        raise ValueError("A planilha não tem uma coluna de data.")
    if shift and not hour_col:
        raise ValueError("A planilha não tem uma coluna de hora para saber o turno.")
    if lines and not line_col:
        raise ValueError("A planilha não diz a linha de produção.")
    ti = _column_index(headers, text_col) - 1 if terms else None
    di = _column_index(headers, date_col) - 1 if (start or end) else None
    hi = _column_index(headers, hour_col) - 1 if shift else None
    li = _column_index(headers, line_col) - 1 if lines else None
    wanted = {str(l).strip().casefold() for l in lines}

    def matches(row_number, values):
        if li is not None and str(values[li]).strip().casefold() not in wanted:
            return False
        if ti is not None:
            value = values[ti]
            if value is None or not any(t in str(value).casefold() for t in terms):
                return False
        if hi is not None:
            value = values[hi]
            if shift_of(value) != shift:
                return False
        if di is not None:
            raw = values[di]
            if raw is None or str(raw).strip() == "":
                return False
            day = parse_date(raw)
            if day is None:
                raise ValueError(f"Data ilegível na linha {row_number}: {raw!r}")
            if (start and day < start) or (end and day > end):
                return False
        return True

    return matches, ti, di, hi


def _as_number(value):
    """The cell as a number, or None. 113 cells of "Produzido" hold the text "0"
    instead of the number, so a plain sum() blows up over the whole sheet."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    return int(text) if re.fullmatch(r"-?\d+", text) else None


def _total(values):
    """(sum of the readable cells, whether every cell could be read).

    A cell nobody could read must never pass for a zero inside a total, so what
    is left out is said out loud instead of silently added as nothing.
    """
    numbers = [n for n in map(_as_number, values) if n is not None]
    return sum(numbers), len(numbers) == len(values)


def _summable(values):
    """Whether adding this column up means anything."""
    numbers = [n for n in map(_as_number, values) if n is not None]
    if not numbers or len(numbers) < 0.8 * len(values):
        return False
    # ponytail: day, shift, line and hour look numeric but are codes. Told apart
    # by being small non-negative integers; if a real 0..31 quantity ever shows
    # up, this becomes a picker on screen.
    return not all(float(n).is_integer() and 0 <= n <= 31 for n in numbers)


def summarize(headers, rows, text_col=None, terms=(), date_col=None, start=None, end=None,
              hour_col=None, shift=None, line_col=None, lines=(), only_with_text=False):
    """Summary of the rows passing the filters.

    Returns the columns that carry content (empty ones are dropped), the matching
    rows, the totals of the summable columns and the date range found. No filter
    at all is valid: it summarises the whole sheet.
    """
    terms = [t.strip().casefold() for t in terms if t.strip()]
    matches, *_ = _matcher(headers, text_col, terms, date_col, start, end, hour_col, shift,
                           line_col, lines)
    ti = _column_index(headers, text_col) - 1 if (only_with_text and text_col) else None
    di = _column_index(headers, date_col) - 1 if date_col else None
    found = [(r, v) for r, v in rows
             if matches(r, v) and (ti is None or str(v[ti] or "").strip())]
    if line_col and line_col in headers:  # keep each line's occurrences together
        by_line = _column_index(headers, line_col) - 1
        found.sort(key=lambda item: str(item[1][by_line]))

    used = [i for i in range(len(headers))
            if headers[i] and any(v[i] not in (None, "", 0) for _, v in found)]
    totals, partial = {}, set()
    for i in used:
        column = [v[i] for _, v in found if v[i] is not None and str(v[i]).strip() != ""]
        if column and _summable(column):
            totals[headers[i]], complete = _total(column)
            if not complete:
                partial.add(headers[i])
    days = sorted(d for d in ((parse_date(v[di]) if di is not None else None) for _, v in found) if d)

    return {
        "matched": found,  # raw rows, for the numbers behind the second tab
        "columns": [headers[i] for i in used],
        "rows": [[v[i] for i in used] for _, v in found],
        "totals": totals,
        "partial": partial,  # totals where some cell could not be read
        "read": len(rows),
        "found": len(found),
        "period": (days[0], days[-1]) if days else None,
    }


def _sheet_path(archive, sheet):
    """xl/worksheets/sheetN.xml of a sheet, via the package rels — openpyxl's
    ws.path answers sheetNone.xml when the workbook was read from disk."""
    book = ET.fromstring(archive.read("xl/workbook.xml"))
    rid = next((s.get(NS_REL + "id") for s in book.iter(NS_MAIN + "sheet") if s.get("name") == sheet), None)
    if rid is None:
        raise ValueError(f"Aba {sheet!r} não encontrada no arquivo.")
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    target = next(r.get("Target") for r in rels if r.get("Id") == rid)
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)


def download(dst, headers, rows, line_col, date_col, hour_col, text_col, terms=()):
    """New .xlsx with just Linha, Data, Hora and Ocorrência of the given rows.

    The occurrence is what the screen shows: with a keyword, only the sentences
    that carry it. A column the sheet does not have comes out blank, never invented.
    """
    terms = [t.strip().casefold() for t in terms if t.strip()]
    at = [_at(headers, c) if c else None for c in (line_col, date_col, hour_col, text_col)]
    book = Workbook()
    sheet = book.active
    sheet.title = "Ocorrências"
    sheet.append(["Linha", "Data", "Hora", "Ocorrência"])
    for _, values in rows:
        line, day, hour, text = (values[i] if i is not None else None for i in at)
        day = parse_date(day) if day is not None else None
        sheet.append([show(line), day, show(hour),
                      sentences_with(text, terms) if text is not None else ""])
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "DD/MM/YYYY"
    for column, width in zip("ABCD", (10, 12, 16, 100)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(dst)
    return len(rows)


def thousands(value):
    """Number the pt-BR way: 976430 -> 976.430."""
    return f"{value:,.0f}".replace(",", ".") if isinstance(value, (int, float)) else show(value)


def show(value):
    """Cell value as the UI shows it."""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y") if (value.hour, value.minute) == (0, 0) \
            else value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value)


def start_dir():
    """Where the file dialog opens: the Windows Downloads folder under WSL, else home."""
    def files(folder):  # system profiles ("Public", "Usuário Padrão") have none
        try:
            return sum(1 for _ in folder.iterdir())
        except OSError:
            return -1

    best = max([Path.home() / "Downloads", *Path("/mnt/c/Users").glob("*/Downloads")], key=files)
    return str(best) if files(best) > 0 else str(Path.home())


# --- Numbers behind the second tab -------------------------------------------
SKU_RE = re.compile(r"\b(\d{7,9})\b")  # the code is the stable part: the wording varies
EQUIPMENT_RE = re.compile(r"^\s*[*•]?\s*([^\n:–-]{3,40}?)\s*[:–-]\s+\S")


def _at(headers, name):
    """Index of a column, or None — the board must survive a sheet without it."""
    return headers.index(name) if name in headers else None


def _values(headers, rows, name):
    i = _at(headers, name)
    return [] if i is None else [v[i] for _, v in rows]


def _add(buckets, key, values, plan_at, made_at):
    """Collect the readable numbers of a bucket; what cannot be read is not a zero."""
    planned, produced = buckets.setdefault(key, ([], []))
    for column, into in ((plan_at, planned), (made_at, produced)):
        number = _as_number(values[column]) if column is not None else None
        if number is not None:
            into.append(number)


def _bucket(collected):
    """(planned, produced) of a bucket — planned is None when no hour of it said one.

    Five hours of February have an empty Plan cell. Summing them as zero would
    paint the bar green, as if production had beaten a target nobody wrote down.
    """
    planned, produced = collected
    return (sum(planned) if planned else None), sum(produced)


def production_by_day(headers, rows):
    """[(day, planned, produced)] — one point per day, in date order."""
    date_at, plan_at, made_at = (_at(headers, n) for n in ("Data", "Plan", "Produzido"))
    if date_at is None:
        return []
    per_day = {}
    for _, values in rows:
        day = parse_date(values[date_at])
        if day is None:
            continue
        _add(per_day, day, values, plan_at, made_at)
    return [(day, *_bucket(per_day[day])) for day in sorted(per_day)]


def production_by_hour(headers, rows):
    """[(hour, planned, produced)] in shift order — what a single day looks like."""
    hour_at, plan_at, made_at = (_at(headers, n) for n in ("Horário", "Plan", "Produzido"))
    if hour_at is None:
        return []
    per_hour = {}
    for _, values in rows:
        hour = str(values[hour_at] or "").strip()
        if not hour:
            continue
        _add(per_hour, hour, values, plan_at, made_at)
    return [(hour, *_bucket(per_hour[hour])) for hour in per_hour]


def production_by_slot(headers, rows):
    """[((day, hour), planned, produced)] — one bar per hour across the whole period.

    A few days plotted per day are four lonely bars in an empty chart; per hour
    they read like the shift panel. The rows already come in day/shift/hour order.
    """
    date_at, hour_at = _at(headers, "Data"), _at(headers, "Horário")
    plan_at, made_at = _at(headers, "Plan"), _at(headers, "Produzido")
    if date_at is None or hour_at is None:
        return []
    per_slot = {}
    for _, values in rows:
        day, hour = parse_date(values[date_at]), str(values[hour_at] or "").strip()
        if day is None or not hour:
            continue
        _add(per_slot, (day, hour), values, plan_at, made_at)
    return [(slot, *_bucket(per_slot[slot])) for slot in per_slot]


def production_by_month(headers, rows):
    """[(month, planned, produced)] — a period longer than a month is unreadable day by day."""
    months = {}
    for day, planned, produced in production_by_day(headers, rows):
        was_planned, was_produced = months.get(day.replace(day=1), (None, 0))
        months[day.replace(day=1)] = (
            None if planned is None and was_planned is None else (was_planned or 0) + (planned or 0),
            was_produced + produced)
    return [(month, *months[month]) for month in sorted(months)]


DURATION_RE = re.compile(
    r"(?:(\d+)\s*h(?:oras?|rs?)?)?\s*(?:(\d+)\s*m(?:inutos?|in)?)?\s*(?:(\d+)\s*s)?", re.I)


def changeover_minutes(text):
    """"31m19s" -> 31.3, "22min" -> 22, "0 Minutos (TROCA DE ORDEM)" -> 0, lixo -> None.

    The field is typed by hand and comes in every spelling the shift leader felt like.
    """
    said = str(text or "").strip()
    if not said:
        return None
    found = DURATION_RE.match(said)
    hours, minutes, seconds = (int(g) if g else 0 for g in found.groups()) if found else (0, 0, 0)
    if not any((hours, minutes, seconds)):
        return 0.0 if re.match(r"^0\b", said) else None
    return hours * 60 + minutes + seconds / 60


def changeover_time_by_day(headers, rows):
    """[(day, minutes of changeover)] — the field belongs to the shift, counted once."""
    date_at, shift_at, swap_at = (_at(headers, n) for n in ("Data", "Turno", "Trocas"))
    if date_at is None or swap_at is None:
        return []
    per_day, seen = {}, set()
    for _, values in rows:
        day = parse_date(values[date_at])
        key = (day, values[shift_at] if shift_at is not None else 0)
        if day is None or key in seen:
            continue
        seen.add(key)
        per_day[day] = per_day.get(day, 0) + (changeover_minutes(values[swap_at]) or 0)
    return [(day, per_day[day]) for day in sorted(per_day)]


def losses_by_equipment(headers, rows, top=10):
    """[(equipment, how many times it shows up)] — "Gasser - alarme..." counts as Gasser.

    One line of the losses cell is one occurrence: the shift log lists them one per
    line, each starting with the machine that failed.
    """
    counted, seen_as = Counter(), {}
    for text in _values(headers, rows, "Principais perdas"):
        for line in str(text or "").splitlines():
            found = EQUIPMENT_RE.match(line)
            if not found:
                continue
            name = " ".join(found.group(1).split())
            counted[name.casefold()] += 1
            seen_as.setdefault(name.casefold(), name)
    return [(seen_as[key], n) for key, n in counted.most_common(top)]


# the event is a line that STARTS with it; "atrasando o change over" in the middle of a
# sentence is context, not another changeover. "chang\w*" covers the typos found in the
# sheet: "change over", "changeover", "CHANGER OVER", "Chang Over".
CHANGEOVER_RE = re.compile(
    r"^\s*[-*•]?\s*(chang\w*\s?over|troca de (formato|variante|sku|produto))", re.I)


def _skus_of(text):
    """[(code, description)] of a SKU cell — a changeover day holds two in one cell."""
    found = []
    for line in str(text or "").splitlines():
        code = SKU_RE.search(line)
        if code:
            found.append((code.group(1), " ".join(line.split())))
    return found


def skus_by_day(headers, rows):
    """[(day, [(code, description)], [(hour, what was written, how long)])] — SKU and changeovers.

    The changeover hour only exists in the shift log ("Change over: troca de
    formato..."), while how long it took sits in the block header (TROCA 1 - TEMPO).
    """
    date_at, sku_at = _at(headers, "Data"), _at(headers, "SKU")
    hour_at, text_at = _at(headers, "Período") or _at(headers, "Horário"), _at(headers, "Principais perdas")
    swap_at, shift_at, file_at = _at(headers, "Trocas"), _at(headers, "Turno"), _at(headers, "Arquivo")
    if date_at is None or sku_at is None:
        return []
    per_day, timed, told, sources = {}, {}, set(), {}
    for _, values in rows:
        day = parse_date(values[date_at])
        if day is None:
            continue
        skus, swaps = per_day.setdefault(day, ([], {}))
        if file_at is not None and values[file_at]:
            sources.setdefault(day, str(values[file_at]))
        if swap_at is not None and str(values[swap_at] or "").strip():
            timed[(day, values[shift_at] if shift_at is not None else 0)] = str(values[swap_at])
        for sku in _skus_of(values[sku_at]):
            if sku[0] not in [code for code, _ in skus]:
                skus.append(sku)
        for line in str(values[text_at] or "" if text_at is not None else "").splitlines():
            if not CHANGEOVER_RE.match(line):
                continue
            # the same changeover is written again in every hour it spans: one entry,
            # every hour it was mentioned in
            said = " ".join(line.split())
            hour = str(values[hour_at]) if hour_at is not None else ""
            hours, _ = swaps.setdefault(said, ([], str(values[swap_at] or "")
                                               if swap_at is not None else ""))
            if hour and hour not in hours:
                hours.append(hour)
            told.add((day, values[shift_at] if shift_at is not None else 0))
    # a changeover the header timed but nobody wrote down still happened: show it
    # with the hour missing instead of dropping it
    for (day, shift), took in timed.items():
        if (day, shift) not in told:
            per_day[day][1][f"troca cronometrada no {shift}º turno, sem relato escrito"] = ([], took)
    return [{"day": day, "file": sources.get(day, ""), "skus": skus,
             "swaps": _in_hour_order([(hours, said, took)
                                      for said, (hours, took) in swaps.items()])}
            for day, (skus, swaps) in sorted(per_day.items())]


def _in_hour_order(swaps):
    """Changeovers as they happened, so the Nth one switches to the Nth SKU."""
    return sorted(swaps, key=lambda swap: swap[0][0] if swap[0] else "zz")


def skus_in(headers, rows):
    """[(code, description, hours)] for the period, busiest first."""
    hours, described = Counter(), {}
    for value in _values(headers, rows, "SKU"):
        found = SKU_RE.search(str(value or ""))
        if not found:
            continue
        code = found.group(1)
        hours[code] += 1
        text = " ".join(str(value).split())
        if len(text) > len(described.get(code, "")):
            described[code] = text
    return [(code, described[code], n) for code, n in hours.most_common()]


def production_summary(headers, rows):
    """Planned x produced for the filtered period, plus what is missing."""
    planned, planned_ok = _total(_values(headers, rows, "Plan"))
    produced, produced_ok = _total(_values(headers, rows, "Produzido"))
    return {"plan": planned, "produzido": produced, "faltou": planned - produced,
            "atingido": produced / planned if planned else None,
            "completo": planned_ok and produced_ok,
            "skus": skus_in(headers, rows), "skus_por_dia": skus_by_day(headers, rows), "dias": production_by_day(headers, rows),
            "horas": production_by_hour(headers, rows),
            "slots": production_by_slot(headers, rows),
            "meses": production_by_month(headers, rows),
            "trocas_por_dia": changeover_time_by_day(headers, rows),
            "equipamentos": losses_by_equipment(headers, rows)}


def merge(datasets):
    """Several files read as one: the rows of each, renumbered, in one dataset.

    Two bundles side by side answer "how do line 1 and line 2 compare", which is
    the whole point of keeping the line as a column.
    """
    kept = [d for d in datasets if d and d["rows"]]
    if not kept:
        raise ValueError("Nada para juntar.")
    headers = kept[0]["headers"]
    for other in kept[1:]:
        if other["headers"] != headers:
            raise ValueError("Só dá para juntar arquivos do mesmo tipo de relatório.")
    rows, number = [], 0
    for one in kept:
        for _, values in one["rows"]:
            number += 1
            rows.append((number, values))
    months = [m for one in kept for m in one.get("months", [])]
    read_months = [when for _, when, _ in months if when and all(when)]
    # a dataset that is itself a merge already lists its sources: keep all of them,
    # otherwise adding a third file erases the names of the first two
    sources = [name for one in kept
               for name in (one.get("sources") or [one.get("path", one.get("sheet", ""))])]
    return {**kept[0], "rows": rows, "months": months,
            "sheet": "{} arquivos".format(len(sources)),
            "month": min(read_months) if read_months else kept[0].get("month"),
            "sources": sources}


# the whole app is dark, after the reference panel chosen for the Produção tab
BOARD_COLOURS = {"bg": "#0b0d10", "card": "#16191e", "grid": "#272c34", "text": "#e8ebef",
                 "muted": "#8b95a3", "hit": "#8ab8f0", "miss": "#c8601a", "plan": "#ffffff",
                 "swap": "#7d8794", "unknown": "#4a5361", "select": "#22344a",
                 "field": "#1f242b", "button": "#242a32", "accent": "#3d7bd9",
                 "ok": "#5cd08a", "error": "#ff6b6b"}


def chart_columns(numbers):
    """([(day, changeover minutes, planned, produced)], monthly) on one shared axis.

    More than 31 days goes monthly: a picket fence of bars reads as nothing.
    """
    monthly = len(numbers["dias"]) > 31
    swaps = {}
    for day, minutes in numbers["trocas_por_dia"]:
        key = day.replace(day=1) if monthly else day
        swaps[key] = swaps.get(key, 0) + minutes
    made = numbers["meses"] if monthly else numbers["dias"]
    return [(when, swaps.get(when, 0), planned, produced)
            for when, planned, produced in made], monthly


def _short(value):
    """50000 -> "50 mil", 1123000 -> "1,1 mi": the way the axis reads aloud."""
    if value >= 1_000_000:
        return f"{value / 1_000_000:.1f} mi".replace(".", ",")
    return f"{value / 1000:.0f} mil" if value >= 1000 else f"{value:.0f}"


def _minutes(value):
    """Axis label for changeover: minutes, or hours once minutes get long."""
    return f"{value / 60:.0f} h" if value >= 600 else f"{value:.0f} min"


def draw_combo_chart(canvas, columns, monthly=False, selected=None, colours=None):
    """Changeover per day on top, production per day in a strip below, one axis.

    The changeover is the headline: bars of minutes. Under it each day's
    production is blue when it reached the plan, orange when it fell short and
    slate when nobody wrote a plan down; the white tick is the plan. Every column
    is tagged "slot<i>" so the mouse finds it; `selected` gets a highlight behind.

    ponytail: tk.Canvas instead of matplotlib, a dependency that triples the .exe
    to draw rectangles. Swap it the day these charts need axes, zoom or export.
    """
    c = {**BOARD_COLOURS, **(colours or {})}
    canvas.delete("all")
    width = max(canvas.winfo_width(), 300)
    height = max(canvas.winfo_height(), 200)
    if not columns:
        canvas.create_text(width / 2, height / 2, text="Sem dados no filtro", fill=c["muted"])
        return
    small = ("TkDefaultFont", 8)
    x = 60
    for name, colour, line in (("Changeover", c["swap"], False), ("Bateu o plano", c["hit"], False),
                               ("Abaixo do plano", c["miss"], False),
                               ("Plano", c["plan"], True), ("Sem plano", c["unknown"], False)):
        if line:
            canvas.create_line(x, 13, x + 14, 13, fill=colour, width=3)
        else:
            canvas.create_rectangle(x, 7, x + 12, 19, fill=colour, width=0)
        x = canvas.bbox(canvas.create_text(x + 18, 13, text=name, anchor="w", fill=c["muted"],
                                           font=small))[2] + 16

    left, right, top = 60, width - 10, 36
    floor = height - 22  # production floor; the day labels sit under it
    swap_floor = top + (floor - top) * 0.6
    strip_top = swap_floor + 26
    step = (right - left) / len(columns)
    swap_ceiling = max(minutes for _, minutes, _, _ in columns) or 1
    made_ceiling = max(max(planned or 0, produced) for _, _, planned, produced in columns) or 1

    def y_swap(value):
        return swap_floor - value / swap_ceiling * (swap_floor - top)

    def y_made(value):
        return floor - value / made_ceiling * (floor - strip_top)

    if selected is not None and 0 <= selected < len(columns):
        canvas.create_rectangle(left + selected * step, top - 6, left + (selected + 1) * step,
                                height - 2, fill=c["select"], width=0)
    for value in (0, swap_ceiling / 2, swap_ceiling):
        canvas.create_line(left, y_swap(value), right, y_swap(value), fill=c["grid"], dash=(2, 3))
        canvas.create_text(left - 6, y_swap(value), text=_minutes(value), anchor="e",
                           fill=c["muted"], font=small)
    canvas.create_text(left + 4, swap_floor + 13, anchor="w", fill=c["muted"],
                       text="Produção por " + ("mês" if monthly else "dia"),
                       font=("TkDefaultFont", 8, "bold"))
    for value in (0, made_ceiling):
        canvas.create_line(left, y_made(value), right, y_made(value), fill=c["grid"], dash=(2, 3))
        canvas.create_text(left - 6, y_made(value), text=_short(value), anchor="e",
                           fill=c["muted"], font=small)
    every = max(1, int(-(-24 // step)))  # labels never collide, however many days
    wide = min(max(step * 0.62, 2), 40)
    for i, (when, minutes, planned, produced) in enumerate(columns):
        x = left + i * step
        middle = x + (step - wide) / 2
        tags = ("slot", f"slot{i}")
        if minutes:
            canvas.create_rectangle(middle, y_swap(minutes), middle + wide, swap_floor,
                                    fill=c["swap"], width=0, tags=tags)
        if produced:
            paint = (c["unknown"] if planned is None
                     else c["hit"] if produced >= planned else c["miss"])
            canvas.create_rectangle(middle, y_made(produced), middle + wide, floor,
                                    fill=paint, width=0, tags=tags)
        if planned:
            canvas.create_line(x + 1, y_made(planned), x + step - 1, y_made(planned),
                               fill=c["plan"], width=2, tags=tags)
        if i % every == 0 or i == selected:
            on = i == selected
            canvas.create_text(x + step / 2, height - 11,
                               text=when.strftime("%m/%y") if monthly else f"{when.day:02d}",
                               fill=c["text"] if on else c["muted"],
                               font=("TkDefaultFont", 8, "bold" if on else "normal"))
        # invisible column so the whole day answers the mouse, not only the bars
        canvas.create_rectangle(x, top, x + step, floor, outline="", fill="", tags=tags)


def draw_rank_chart(canvas, items, colour="#c62828", highlight=None, ink="black",
                    quiet="#444", strong="#8e1b1b"):
    """Horizontal bars: how many times each machine shows up in the period.

    Each bar is tagged "bar<i>" so the screen can tell which machine the mouse is
    over; `highlight` is the index to emphasise.
    """
    canvas.delete("all")
    width = max(canvas.winfo_width(), 200)
    height = max(canvas.winfo_height(), 120)
    if not items:
        canvas.create_text(width / 2, height / 2, text="Sem ocorrências no filtro", fill="#777")
        return
    label_width, step = 150, min(22, height / max(len(items), 1))
    ceiling = max(n for _, n in items) or 1
    for i, (name, count) in enumerate(items):
        y = i * step + step / 2
        tag = f"bar{i}"
        on = i == highlight
        canvas.create_text(label_width - 6, y, text=name[:26], anchor="e", tags=("bar", tag),
                           fill=ink, font=("TkDefaultFont", 8, "bold" if on else "normal"))
        long = (count / ceiling) * (width - label_width - 46)
        canvas.create_rectangle(label_width, y - step / 3, label_width + long, y + step / 3,
                                fill=strong if on else colour, width=0, tags=("bar", tag))
        canvas.create_text(label_width + long + 6, y, text=str(count), anchor="w",
                           font=("TkDefaultFont", 8), fill=quiet, tags=("bar", tag))
        # invisible strip so the whole row answers the mouse, not just the bar
        canvas.create_rectangle(0, y - step / 2, width, y + step / 2, outline="", fill="",
                                tags=("bar", tag))


def tag_index(canvas, x, y, prefix):
    """Index of the "<prefix><i>" item nearest to (x, y), or None.

    By point, not by Tk's "current" item: the same lookup then works for a real
    mouse and for a generated event, which is what the tests use.
    """
    for item in canvas.find_closest(x, y):
        for tag in canvas.gettags(item):
            if tag.startswith(prefix) and tag[len(prefix):].isdigit():
                return int(tag[len(prefix):])
    return None


def draw_tip(canvas, x, y, lines):
    """Small box next to the cursor with the exact numbers of the bar under it."""
    canvas.delete("tip")
    pad, line_height = 6, 14
    width = max(len(line) for line in lines) * 6.4 + pad * 2
    height = len(lines) * line_height + pad * 2
    left = min(x + 12, max(canvas.winfo_width(), 200) - width - 4)
    top = max(y - height - 10, 2)
    canvas.create_rectangle(left, top, left + width, top + height, fill="#333a45", outline="",
                            tags="tip")
    for i, line in enumerate(lines):
        canvas.create_text(left + pad, top + pad + i * line_height, text=line, anchor="nw",
                           fill="white", font=("TkDefaultFont", 8), tags="tip")


def draw_chips(frame, chosen, on_remove):
    """Redraw the picked lines as chips, each with an x that takes it out again.

    Module level so a test can build it on a real Tk root and click the x.
    """
    import tkinter as tk
    from tkinter import ttk

    for child in frame.winfo_children():
        child.destroy()
    if not chosen:
        ttk.Label(frame, text="todas as linhas", foreground=BOARD_COLOURS["muted"],
                  background=BOARD_COLOURS["card"]).pack(side="left")
        return []
    made = []
    for name in chosen:
        tone = BOARD_COLOURS["select"]
        chip = tk.Frame(frame, background=tone, padx=8, pady=2)
        chip.pack(side="left", padx=(0, 4))
        tk.Label(chip, text=name, background=tone, foreground=BOARD_COLOURS["text"]).pack(side="left")
        close = tk.Button(chip, text="✕", background=tone, activebackground=BOARD_COLOURS["grid"],
                          relief="flat", borderwidth=0, highlightthickness=0, padx=3, pady=0,
                          cursor="hand2", foreground=BOARD_COLOURS["muted"],
                          activeforeground=BOARD_COLOURS["text"],
                          command=lambda n=name: on_remove(n))
        close.pack(side="left")
        made.append(close)
    return made


def main(path=None):
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title(f"Resumo de ocorrências — build {BUILD}")
    dark = BOARD_COLOURS
    root.configure(background=dark["bg"])
    root.geometry("1280x900")
    try:  # the board needs the room: open maximised where the platform knows how
        root.state("zoomed")
    except tk.TclError:
        pass
    style = ttk.Style()
    style.theme_use("clam")  # the native themes ignore most colours; clam takes them all
    style.configure(".", background=dark["bg"], foreground=dark["text"],
                    fieldbackground=dark["field"], bordercolor=dark["grid"],
                    lightcolor=dark["grid"], darkcolor=dark["grid"], troughcolor=dark["bg"],
                    insertcolor=dark["text"], selectbackground=dark["select"],
                    selectforeground=dark["text"], arrowcolor=dark["muted"],
                    font=("TkDefaultFont", 10))
    style.map(".", foreground=[("disabled", dark["muted"])],
              fieldbackground=[("disabled", dark["card"])])
    style.configure("Card.TFrame", background=dark["card"])
    style.configure("Card.TLabel", background=dark["card"])
    style.configure("Field.TLabel", background=dark["card"], foreground=dark["muted"],
                    font=("TkDefaultFont", 9, "bold"))
    style.configure("Hint.TLabel", background=dark["card"], foreground=dark["muted"],
                    font=("TkDefaultFont", 9))
    style.configure("File.TLabel", background=dark["card"], font=("TkDefaultFont", 12, "bold"))
    style.configure("Period.TLabel", font=("TkDefaultFont", 12, "bold"))
    style.configure("Brand.TLabel", font=("TkDefaultFont", 18, "bold"))
    style.configure("Tagline.TLabel", foreground=dark["muted"], font=("TkDefaultFont", 10))
    style.configure("TEntry", padding=(8, 6))
    style.configure("TCombobox", padding=(8, 6))
    # a readonly combobox paints its whole field with the selection colour once it
    # has focus: keep it the field colour, or every picker looks like a blue block
    style.map("TCombobox", fieldbackground=[("readonly", dark["field"]),
                                            ("disabled", dark["card"])],
              selectbackground=[("readonly", dark["field"])],
              selectforeground=[("readonly", dark["text"])])
    root.option_add("*TCombobox*Listbox.background", dark["field"])
    root.option_add("*TCombobox*Listbox.foreground", dark["text"])
    root.option_add("*TCombobox*Listbox.selectBackground", dark["select"])
    root.option_add("*TCombobox*Listbox.selectForeground", dark["text"])
    style.configure("TButton", background=dark["button"], padding=(12, 6), borderwidth=1,
                    width=-8)
    style.map("TButton", background=[("disabled", dark["card"]), ("active", dark["grid"])])
    style.configure("Accent.TButton", background=dark["accent"], foreground="white",
                    bordercolor=dark["accent"], font=("TkDefaultFont", 10, "bold"))
    style.map("Accent.TButton", background=[("disabled", dark["card"]), ("active", "#5b93e6")])
    style.configure("TNotebook", background=dark["bg"], borderwidth=0, bordercolor=dark["grid"],
                    tabmargins=(0, 4, 0, 0))
    style.configure("TNotebook.Tab", background=dark["bg"], foreground=dark["muted"],
                    padding=(16, 8), borderwidth=0, font=("TkDefaultFont", 10, "bold"))
    style.map("TNotebook.Tab", background=[("selected", dark["card"])],
              foreground=[("selected", dark["text"])])
    style.configure("TScrollbar", background=dark["button"], bordercolor=dark["bg"],
                    arrowcolor=dark["muted"])
    style.configure("Treeview", background=dark["field"], fieldbackground=dark["field"],
                    foreground=dark["text"], rowheight=26, bordercolor=dark["grid"])
    style.map("Treeview", background=[("selected", dark["select"])],
              foreground=[("selected", dark["text"])])
    style.configure("Treeview.Heading", background=dark["button"], foreground=dark["muted"],
                    font=("TkDefaultFont", 9, "bold"), padding=(6, 4))
    book = {"path": "", "headers": [], "rows": [], "date": "", "hour": "", "text": "",
            "whole": None, "lines": [], "chosen": [], "admin": None}

    source, detected = tk.StringVar(value="Nenhuma planilha aberta"), tk.StringVar()
    start_date, end_date, hour, keyword = (tk.StringVar() for _ in range(4))
    recent, who = tk.StringVar(), tk.StringVar()
    status = tk.StringVar(value="Carregando a base salva...")
    summary = tk.StringVar(value="—")

    def panel(parent):
        """A card on the black background: one bordered block per job."""
        box = tk.Frame(parent, background=dark["card"], padx=14, pady=12,
                       highlightthickness=1, highlightbackground=dark["grid"])
        return box

    def field(parent, caption, widget, column, span=1, hint=None):
        """Caption above the control, the way a form is read: label, then input."""
        ttk.Label(parent, text=caption.upper(), style="Field.TLabel").grid(
            row=0, column=column, columnspan=span, sticky="w", pady=(0, 4))
        widget.grid(row=1, column=column, columnspan=span, sticky="we", padx=(0, 12))
        if hint:
            ttk.Label(parent, text=hint, style="Hint.TLabel").grid(
                row=2, column=column, columnspan=span, sticky="w", pady=(4, 0))

    # the name of the app first, above everything else
    ttk.Label(root, text="Oráculo Digital", style="Brand.TLabel",
              padding=(14, 10, 14, 0)).pack(fill="x")
    ttk.Label(root, text="enxergando oportunidades", style="Tagline.TLabel",
              padding=(14, 0, 14, 0)).pack(fill="x")

    # file bar: what is open, and — for the admin — how to open or feed something else
    files = panel(root)
    files.pack(fill="x", padx=10, pady=(10, 0))
    file_side = ttk.Frame(files, style="Card.TFrame")
    file_side.grid(row=0, column=0, rowspan=2, sticky="w", padx=(0, 16))
    file_buttons = ttk.Frame(file_side, style="Card.TFrame")
    file_buttons.pack(anchor="w")
    open_button = ttk.Button(file_buttons, text="Abrir...", style="Accent.TButton",
                             command=lambda: open_file())
    add_button = ttk.Button(file_buttons, text="+ Adicionar", command=lambda: open_file(add=True))
    api_button = ttk.Button(file_buttons, text="Adicionar API", command=lambda: open_api())
    login_button = ttk.Button(file_buttons, text="Unilever", command=lambda: toggle_login())
    login_button.pack(side="left")
    ttk.Label(file_side, textvariable=who, style="Hint.TLabel").pack(anchor="w", pady=(4, 0))
    ttk.Label(files, textvariable=source, style="File.TLabel").grid(row=0, column=1, sticky="w")
    # what the app worked out on its own: without this the filters look like magic
    ttk.Label(files, textvariable=detected, style="Hint.TLabel").grid(
        row=1, column=1, columnspan=2, sticky="w")
    files.columnconfigure(1, weight=1)
    reopen_bar = ttk.Frame(files, style="Card.TFrame")
    reopen_bar.grid(row=0, column=2, sticky="e", pady=(0, 4))
    ttk.Label(reopen_bar, text="HISTÓRICO", style="Field.TLabel").pack(side="left", padx=(0, 8))
    recent_box = ttk.Combobox(reopen_bar, textvariable=recent, state="readonly", width=26)
    recent_box.pack(side="left")
    ttk.Button(reopen_bar, text="Reabrir", command=lambda: reopen()).pack(side="left", padx=(8, 0))
    ttk.Button(reopen_bar, text="✕", width=3, command=lambda: drop_history()).pack(side="left",
                                                                                  padx=(4, 0))

    # filters: one row of fields, the keyword under them, the actions on the right
    filters_box = panel(root)
    filters_box.pack(fill="x", padx=10, pady=(10, 0))
    fields = ttk.Frame(filters_box, style="Card.TFrame")
    fields.pack(side="left", fill="x", expand=True)
    first = ttk.Entry(fields, textvariable=start_date, width=14, state="disabled")
    last = ttk.Entry(fields, textvariable=end_date, width=14, state="disabled")
    hour_box = ttk.Combobox(fields, textvariable=hour, width=24, state="disabled",
                            values=["todos"] + list(SHIFT_LABELS.values()))
    hour.set("todos")
    line_choice = tk.StringVar()
    line_box = ttk.Combobox(fields, textvariable=line_choice, state="disabled", width=16)
    field(fields, "Data inicial", first, 0, hint="dd/mm/aaaa ou só o dia")
    field(fields, "Data final", last, 1)
    field(fields, "Turno", hour_box, 2)
    field(fields, "Linhas", line_box, 3)
    chips = ttk.Frame(fields, style="Card.TFrame")
    chips.grid(row=2, column=3, sticky="w", pady=(4, 0))
    words = ttk.Entry(fields, textvariable=keyword, state="disabled")
    field(fields, "Palavra-chave", words, 4, hint="separe por vírgula para buscar qualquer uma"
                                                   " · Enter filtra")
    fields.columnconfigure(4, weight=1)
    actions = ttk.Frame(filters_box, style="Card.TFrame")
    actions.pack(side="right", padx=(4, 0))
    filter_button = ttk.Button(actions, text="Filtrar", style="Accent.TButton",
                               command=lambda: refresh(), state="disabled")
    filter_button.pack(fill="x", pady=(0, 6))
    download_button = ttk.Button(actions, text="Baixar Dados", command=lambda: download_file(),
                                 state="disabled")
    download_button.pack(fill="x")

    headline = ttk.Label(root, textvariable=summary, style="Period.TLabel", padding=(14, 10, 14, 2))
    headline.pack(fill="x")
    headline.bind("<Configure>", lambda e: headline.configure(wraplength=e.width - 24))

    tabs = ttk.Notebook(root)  # packed after the status line, so that line never gets cut

    frame = tk.Frame(tabs, background=dark["bg"], padx=10, pady=10)
    tabs.add(frame, text="  Ocorrências  ")
    listing = tk.Text(frame, wrap="word", padx=16, pady=12, spacing3=4, state="disabled",
                      background=dark["card"], foreground=dark["text"], relief="flat",
                      highlightthickness=1, highlightbackground=dark["grid"],
                      font=("TkDefaultFont", 10), cursor="arrow")
    bar = ttk.Scrollbar(frame, orient="vertical", command=listing.yview)
    listing.configure(yscrollcommand=bar.set)
    listing.pack(side="left", fill="both", expand=True)
    bar.pack(side="right", fill="y")
    listing.tag_config("when", font=("TkDefaultFont", 10, "bold"), foreground=dark["hit"])
    listing.tag_config("none", foreground=dark["muted"])
    listing.tag_config("hit", background="#6b5310", foreground="#fff3c4")
    listing.tag_config("line", font=("TkDefaultFont", 12, "bold"), foreground=dark["text"],
                       spacing1=10, spacing3=6)

    board = tk.Frame(tabs, background=dark["bg"], padx=10, pady=10)
    tabs.insert(0, board, text="  Produção  ")  # the main screen: it opens here
    tabs.select(board)

    def card(parent, title):
        """Dark panel with a bold title, like the rest of the board."""
        box = tk.Frame(parent, background=dark["card"], padx=14, pady=10,
                       highlightthickness=1, highlightbackground=dark["grid"])
        heading = tk.Label(box, background=dark["card"], foreground=dark["text"],
                           font=("TkDefaultFont", 11, "bold"), anchor="w",
                           **({"textvariable": title} if isinstance(title, tk.StringVar)
                              else {"text": title}))
        heading.pack(fill="x", pady=(0, 6))
        return box

    tiles = tk.Frame(board, background=dark["bg"])
    tiles.pack(fill="x")
    tile_values = {}
    for i, (key, caption) in enumerate((("plan", "Planejado"), ("produzido", "Produzido"),
                                        ("faltou", "Faltou"), ("atingido", "Atingido"),
                                        ("ocorrencias", "Linhas de ocorrência"))):
        cell = tk.Frame(tiles, background=dark["card"], padx=14, pady=8,
                        highlightthickness=1, highlightbackground=dark["grid"])
        cell.grid(row=0, column=i, sticky="we", padx=(0 if i == 0 else 8, 0))
        tiles.columnconfigure(i, weight=1)
        tk.Label(cell, text=caption, background=dark["card"],
                 foreground=dark["muted"]).pack(anchor="w")
        tile_values[key] = tk.Label(cell, text="—", background=dark["card"],
                                    foreground=dark["text"], font=("TkDefaultFont", 15, "bold"))
        tile_values[key].pack(anchor="w")

    # machines first; under them the changeover chart and the SKU list side by side
    rank_card = card(board, "Ocorrências por equipamento no período")
    rank_card.pack(fill="x", pady=(10, 0))
    rank_chart = tk.Canvas(rank_card, height=150, background=dark["card"], highlightthickness=0)
    rank_chart.pack(fill="x")

    pair = tk.Frame(board, background=dark["bg"])
    pair.pack(fill="both", expand=True, pady=(10, 0))
    pair.columnconfigure(0, weight=1)
    pair.rowconfigure(0, weight=1)
    chart_title = tk.StringVar(value="Tempo de troca (changeover) e produção por dia")
    chart_card = card(pair, chart_title)
    chart_card.grid(row=0, column=0, sticky="nsew")
    # the hint is packed first: on a short screen the chart shrinks, the hint stays
    tk.Label(chart_card, background=dark["card"], foreground=dark["muted"], anchor="w",
             text="Clique num dia para ver o SKU e as trocas dele ao lado · "
                  "duplo clique filtra só esse dia").pack(side="bottom", fill="x", pady=(6, 0))
    day_chart = tk.Canvas(chart_card, height=240, background=dark["card"], highlightthickness=0)
    day_chart.pack(fill="both", expand=True)

    sku_title = tk.StringVar(value="SKU e changeover por dia")
    sku_card = card(pair, sku_title)
    sku_card.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
    sku_list = tk.Text(sku_card, wrap="word", width=52, height=14, padx=4, pady=2,
                       font=("TkDefaultFont", 9),
                       state="disabled", background=dark["card"], foreground=dark["text"],
                       relief="flat", highlightthickness=0, cursor="arrow")
    sku_list.pack(fill="both", expand=True)
    sku_list.tag_config("day", font=("TkDefaultFont", 10, "bold"), spacing1=6)
    sku_list.tag_config("picked", background=dark["select"])
    sku_list.tag_config("code", font=("TkDefaultFont", 9, "bold"), foreground=dark["hit"])
    sku_list.tag_config("swap", foreground="#e0a060")
    sku_list.tag_config("quiet", foreground=dark["muted"])

    status_line = ttk.Label(root, textvariable=status, padding=(14, 6), foreground=dark["muted"])
    status_line.pack(side="bottom", fill="x")
    tabs.pack(fill="both", expand=True, padx=10, pady=(0, 4))

    def say(message, tone=BOARD_COLOURS["muted"]):
        status.set(message)
        status_line.configure(foreground=tone)

    def show_history(entries=None):
        book["history"] = entries if entries is not None else history()
        recent_box["values"] = [
            "{}  ·  {}  ·  {} linhas{}".format(
                Path(e["path"]).name, e["when"][:16].replace("T", " "), e.get("rows", "?"),
                "  ·  {:02d}/{}".format(e["month"][1], e["month"][0]) if e.get("month") else "")
            for e in book["history"]]

    def reopen():
        picked = recent_box.current()
        if picked >= 0:
            entry = book["history"][picked]
            month = tuple(entry["month"]) if entry.get("month") else None
            open_file(entry["path"], month)

    def drop_history():
        """The x next to Reabrir: takes the chosen file out of the history."""
        picked = recent_box.current()
        if picked < 0:
            show_history(forget_all())
            say("Histórico apagado.")
            return
        gone = book["history"][picked]
        show_history(forget(gone["path"]))
        recent.set("")
        say(f"{Path(gone['path']).name} saiu do histórico.")

    def open_file(chosen=None, month=None, add=False):
        if book.get("reading"):
            say("Espere a leitura em andamento terminar.", dark["error"])
            return
        if add and not book["rows"]:
            say("Abra o primeiro arquivo antes de adicionar outro.", dark["error"])
            return
        if not chosen:
            opened = Path(book["path"]).parent if book["path"] else start_dir()
            chosen = filedialog.askopenfilename(
                filetypes=[("Excel ou pacote", "*.xlsx *.zip"), ("Excel", "*.xlsx"),
                           ("Pacote de relatórios", "*.zip")], initialdir=str(opened))
        if not chosen:
            return
        if book["rows"] and not add:
            # both buttons end up here; asking removes any doubt about which one was hit
            somar = messagebox.askyesnocancel(
                "Já há dados abertos",
                "Somar {} ao que já está aberto?\n\n"
                "Sim = somar     Não = substituir     Cancelar = deixar como está".format(
                    Path(chosen).name))
            if somar is None:
                say("Cancelado — nada mudou.")
                return
            add = somar
        if not Path(chosen).exists():
            say(f"{chosen} não está mais nesse caminho.", dark["error"])
            return
        source.set(f"Lendo {Path(chosen).name}...")
        say("Lendo em segundo plano — a janela continua respondendo.")
        book["reading"], book["adding"] = True, add
        for widget in (first, last, words, filter_button, download_button):
            widget.state(["disabled"])
        news = queue.Queue()

        def work():  # off the UI thread: Tk is touched only by the poller below
            try:
                news.put(("ok", load(chosen, month,
                                     progress=lambda i, n, name: news.put(("step", (i, n, name))))))
            except Exception as error:
                news.put(("erro", error))

        threading.Thread(target=work, daemon=True).start()

        def poll():
            try:
                kind, payload = news.get_nowait()
            except queue.Empty:
                root.after(100, poll)
                return
            if kind == "step":
                done, total, name = payload
                source.set(f"Lendo {done}/{total}: {name}")
                root.after(50, poll)
                return
            book["reading"] = False
            for widget in (first, last, words, filter_button, download_button):
                widget.state(["!disabled"])
            if kind == "erro":
                source.set("Nenhuma planilha aberta")
                say(f"Erro ao ler: {payload}", dark["error"])
                return
            finish(payload)

        root.after(100, poll)

    def finish(data):
        chosen = data["path"] if isinstance(data, dict) and data.get("path") else book["path"]
        add = book.get("adding", False)
        if add:
            try:
                data = merge([book["data"], data])
            except ValueError as error:
                say(f"Erro ao juntar: {error}", dark["error"])
                return
        book["just_loaded"] = (
            "Base salva: o que o admin alimentou pelo botão API." if data.get("base")
            else "Somado ao que já estava aberto: {} arquivos.".format(len(data.get("sources", [])))
            if add else "Abriu do zero — o que estava aberto foi substituído. Use + Adicionar"
            " para somar.")
        book["data"] = data
        book.update(path=data.get("path", chosen), headers=data["headers"],
                    rows=data["rows"], kind=data["kind"])
        picked = data["columns"]
        book["date"], book["hour"] = picked["date"], picked["hour"]
        book["label"], book["text"] = picked["label"], picked["text"]
        book["line"] = picked.get("line", "")
        book["month"] = data["month"]
        source.set("{}   ·   {} linhas".format(
            "Base salva" if data.get("base")
            else " + ".join(Path(p).name for p in data["sources"]) if data.get("sources")
            else Path(chosen).name, len(data["rows"])))
        if data.get("months"):  # zip: several workbooks, each with its own month
            lidos = [f"{when[1]:02d}/{when[0]}" for _, when, _ in data["months"] if when and all(when)]
            perdidos = [name for name, when, _ in data["months"] if not (when and all(when))]
            origem = "{} planilhas · {}".format(len(data["months"]), ", ".join(sorted(set(lidos))))
            if perdidos:
                origem += f"   ·   SEM mês: {', '.join(perdidos)}"
        elif data["kind"] == "turno":
            day_at = data["headers"].index("Dia")
            origem = "lido das {} abas por dia".format(len({r[day_at] for _, r in data["rows"]}))
        else:
            origem = f"aba {data['sheet']}"
        vindo = data.get("month_from")
        if vindo:
            sobre_o_mes = f"   ·   mês veio da aba {vindo}"
        elif data.get("month_guessed"):
            sobre_o_mes = "   ·   NENHUMA aba diz o mês — confira o campo"
        else:
            sobre_o_mes = ""
        detected.set("{}   ·   Data: {}   ·   Hora: {}   ·   Ocorrência: {}{}".format(
            origem, book["date"] or "não encontrada", book["label"] or "não encontrada",
            book["text"] or "não encontrada", sobre_o_mes))
        book["whole"] = summarize(book["headers"], book["rows"], date_col=book["date"],
                                  text_col=book["text"])
        book["whole"]["with_text"] = summarize(book["headers"], book["rows"],
                                               text_col=book["text"], only_with_text=True)["found"]
        book["lines"] = (distinct(book["headers"], book["rows"], picked["line"])
                         if picked.get("line") else [])
        book["chosen"] = []
        line_box.configure(state="readonly" if book["lines"] else "disabled")
        redraw()
        for field in (start_date, end_date, keyword):
            field.set("")
        hour.set("todos")
        for widget in (first, last, words, filter_button, download_button):
            widget.state(["!disabled"])
        hour_box.configure(state="readonly" if book["hour"] else "disabled")
        if not add and not data.get("base"):
            show_history(remember(chosen, {"rows": len(data["rows"]), "month": data["month"]}))
        refresh()

    def clear_screen():
        """Nothing to show: every number back to a dash, no stale chart left up."""
        book.update(rows=[], headers=[], data=None, board=None, columns=[], path="", lines=[],
                    chosen=[])
        for text in (listing, sku_list):
            text.configure(state="normal")
            text.delete("1.0", "end")
            text.configure(state="disabled")
        for label in tile_values.values():
            label.configure(text="—", foreground=dark["text"])
        summary.set("—")
        detected.set("")
        draw_combo_chart(day_chart, [])
        draw_rank_chart(rank_chart, [])
        for widget in (first, last, words, filter_button, download_button):
            widget.state(["disabled"])
        hour_box.configure(state="disabled")
        line_box.configure(state="disabled")
        redraw()

    def show_base(announce=None, keep_filters=False):
        """Put the saved base on screen — what everyone sees when the app opens.

        An automatic refresh keeps the dates, turno, lines and keyword the person
        was looking at: a change in SharePoint must not undo their filter.
        """
        data = base_load() or {}
        missing = data.get("missing") or []
        if not data.get("rows"):
            clear_screen()
            source.set("Base vazia")
            if missing:
                say("Nada da base abriu: " + ", ".join(missing), dark["error"])
            else:
                say("A base está vazia — o admin entra em Unilever e alimenta em Adicionar API.")
            return
        kept = ((start_date.get(), end_date.get(), keyword.get(), hour.get(),
                 list(book.get("chosen") or [])) if keep_filters and book.get("rows") else None)
        book["adding"] = False
        finish(data)
        if kept:
            for field, value in zip((start_date, end_date, keyword, hour), kept):
                field.set(value)
            book["chosen"] = [line for line in kept[4] if line in book["lines"]]
            redraw()
            refresh()
        if missing:  # smaller totals must not pass for whole ones
            say("Não abriram da base, e os totais estão sem eles: " + ", ".join(missing),
                dark["error"])
        elif announce:
            say(announce, dark["ok"])

    def centre(dialog, size=None):
        """Open a dialog over the middle of the window, not wherever the system drops it."""
        dialog.update_idletasks()
        width, height = size or (dialog.winfo_reqwidth(), dialog.winfo_reqheight())
        x = max(root.winfo_rootx() + (root.winfo_width() - width) // 2, 0)
        y = max(root.winfo_rooty() + (root.winfo_height() - height) // 3, 0)
        dialog.geometry(f"{width}x{height}+{x}+{y}" if size else f"+{x}+{y}")

    def apply_role():
        """Viewer: the Unilever button. Admin: Abrir, + Adicionar, Deslogar, Adicionar API."""
        admin = book.get("admin")
        for button in (open_button, add_button):
            if admin:
                button.pack(side="left", padx=(0, 8), before=login_button)
            else:
                button.pack_forget()
        if admin:
            api_button.pack(side="left", padx=(8, 0), after=login_button)
            reopen_bar.grid()
        else:
            api_button.pack_forget()
            reopen_bar.grid_remove()
            if book.get("api") and book["api"].winfo_exists():
                book["api"].destroy()  # an admin window must not outlive the admin
        login_button.configure(text="Deslogar" if admin else "Unilever")
        who.set(f"admin: {admin}" if admin else "")

    def toggle_login():
        """Unilever opens the login; once in, the same button is Deslogar."""
        if book.get("admin"):
            book["admin"] = None
            apply_role()
            say("Deslogado.")
            return
        dialog = tk.Toplevel(root, background=dark["card"], padx=24, pady=20)
        dialog.title("Unilever — entrar")
        dialog.transient(root)
        dialog.resizable(False, False)
        user, secret, problem = tk.StringVar(), tk.StringVar(), tk.StringVar()
        ttk.Label(dialog, text="LOGIN", style="Field.TLabel").pack(anchor="w")
        user_entry = ttk.Entry(dialog, textvariable=user, width=34)
        user_entry.pack(fill="x", pady=(4, 12))
        ttk.Label(dialog, text="SENHA", style="Field.TLabel").pack(anchor="w")
        secret_entry = ttk.Entry(dialog, textvariable=secret, show="•", width=34)
        secret_entry.pack(fill="x", pady=(4, 6))
        ttk.Label(dialog, textvariable=problem, style="Card.TLabel",
                  foreground=dark["error"]).pack(anchor="w")

        def enter(*_):
            if not admin_file().exists():
                problem.set("Este .exe foi gerado sem admin configurado.")
            elif is_admin(user.get(), secret.get()):
                book["admin"] = user.get().strip()
                dialog.destroy()
                apply_role()
                say(f"Logado como {book['admin']}: Adicionar API alimenta a base.", dark["ok"])
            else:
                secret.set("")
                problem.set("Login ou senha não conferem.")
                secret_entry.focus_set()

        ttk.Button(dialog, text="Entrar", style="Accent.TButton", command=enter).pack(
            fill="x", pady=(10, 0))
        dialog.bind("<Return>", enter)
        dialog.bind("<Escape>", lambda _: dialog.destroy())
        centre(dialog)
        user_entry.focus_set()
        dialog.grab_set()

    def sync_now():
        """Check the live sources off the UI thread; refresh the screen if any changed."""
        if book.get("syncing") or not load_sources()["sources"]:
            return
        book["syncing"] = True
        news = queue.Queue()

        def work():
            try:
                news.put(("ok", sync_sources()))
            except Exception as error:
                news.put(("erro", error))

        threading.Thread(target=work, daemon=True).start()

        def poll():
            try:
                kind, payload = news.get_nowait()
            except queue.Empty:
                root.after(200, poll)
                return
            book["syncing"] = False
            if book.get("api_refresh"):
                book["api_refresh"]()
            if kind == "erro":
                say(f"Erro ao atualizar as fontes: {payload}", dark["error"])
                return
            if payload:
                said = "Atualizado às {}: {}.".format(datetime.now().strftime("%H:%M"), ", ".join(
                    label + (f" (modificado por {who})" if who else "") for label, who in payload))
                if not book.get("rows") or (book.get("data") or {}).get("base"):
                    show_base(said, keep_filters=True)
                else:  # the admin opened a file by hand: say it, do not swap the screen
                    say(said + " Mostre a base em Adicionar API para ver.", dark["ok"])
            broken = [s for s in load_sources()["sources"] if s.get("problem")]
            problems = {(s["id"], s["problem"]) for s in broken}
            if problems and problems != book.get("problems"):  # say it once, not every minute
                say(f"Fonte {broken[0]['label']}: {broken[0]['problem']}", dark["error"])
            book["problems"] = problems

        root.after(200, poll)

    def keep_syncing():
        sync_now()
        root.after(SYNC_EVERY_MS, keep_syncing)

    def open_api():
        """Admin only: .zip, the Graph token, live sources per month, and the base itself."""
        if not book.get("admin"):
            return
        if book.get("api") and book["api"].winfo_exists():
            book["api"].lift()
            return
        dialog = book["api"] = tk.Toplevel(root, background=dark["card"], padx=20, pady=16)
        dialog.title("Adicionar API")
        dialog.transient(root)
        pages = ttk.Notebook(dialog)
        pages.pack(fill="both", expand=True)
        live = tk.Frame(pages, background=dark["card"], padx=4, pady=12)
        stored = tk.Frame(pages, background=dark["card"], padx=4, pady=12)
        pages.add(live, text="  Fontes em tempo real  ")
        pages.add(stored, text="  Na base  ")
        state = load_sources()
        token, label, key, note = (tk.StringVar(value=state["token"]), tk.StringVar(),
                                   tk.StringVar(), tk.StringVar())

        zip_row = ttk.Frame(live, style="Card.TFrame")
        zip_row.pack(fill="x")
        send_button = ttk.Button(zip_row, text="Enviar .zip...", command=lambda: feed_file())
        send_button.pack(side="left")
        ttk.Label(zip_row, style="Hint.TLabel", text="importa uma vez · a mesma linha no mesmo mês"
                  " substitui a anterior").pack(side="left", padx=(12, 0))

        ttk.Label(live, text="TOKEN / CÓDIGO (MICROSOFT GRAPH)", style="Field.TLabel").pack(
            anchor="w", pady=(16, 4))
        token_row = ttk.Frame(live, style="Card.TFrame")
        token_row.pack(fill="x")
        ttk.Entry(token_row, textvariable=token, show="•").pack(side="left", fill="x", expand=True)
        ttk.Button(token_row, text="Salvar token", command=lambda: save_token()).pack(
            side="left", padx=(8, 0))
        ttk.Label(live, style="Hint.TLabel", wraplength=930, text="O token de acesso do Graph vale"
                  " cerca de 1 hora. Sem token, a chave precisa ser um link que baixe sem login ou"
                  " um arquivo da pasta que o OneDrive sincroniza.").pack(anchor="w", pady=(4, 0))

        add_row = ttk.Frame(live, style="Card.TFrame")
        add_row.pack(fill="x", pady=(16, 0))
        ttk.Label(add_row, text="NOME", style="Field.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(add_row, text="CHAVE", style="Field.TLabel").grid(row=0, column=1, sticky="w")
        ttk.Combobox(add_row, textvariable=label, width=22,
                     values=[str(date.today().year), *MONTH_NAMES]).grid(
            row=1, column=0, sticky="w", padx=(0, 8), pady=(4, 0))
        ttk.Entry(add_row, textvariable=key).grid(row=1, column=1, sticky="we", pady=(4, 0))
        probe_button = ttk.Button(add_row, text="Testar conexão", command=lambda: probe())
        probe_button.grid(row=1, column=2, padx=(8, 0), pady=(4, 0))
        ttk.Button(add_row, text="Adicionar fonte", style="Accent.TButton",
                   command=lambda: add()).grid(row=1, column=3, padx=(8, 0), pady=(4, 0))
        add_row.columnconfigure(1, weight=1)
        ttk.Label(live, style="Hint.TLabel", wraplength=930, text="Chave = drives/<id>/items/<id>"
                  " de um arquivo ou de uma pasta (o ano, o mês — o app percorre tudo que tiver"
                  " .xlsx dentro), a URL do Graph com ela (/content ou /children também servem),"
                  " o link do SharePoint, ou o arquivo/pasta que o OneDrive sincroniza.").pack(
            anchor="w", pady=(4, 10))

        sources_table = ttk.Treeview(live, show="headings", selectmode="browse", height=6,
                                     columns=("mes", "chave", "quando", "por", "situacao"))
        for column, caption, width in (("mes", "Nome", 140), ("chave", "Chave", 220),
                                       ("quando", "Atualizado", 110), ("por", "Por", 120),
                                       ("situacao", "Situação", 300)):
            sources_table.heading(column, text=caption, anchor="w")
            sources_table.column(column, width=width, anchor="w",
                                 stretch=column in ("chave", "situacao"))
        sources_table.pack(fill="both", expand=True)
        live_bottom = ttk.Frame(live, style="Card.TFrame")
        live_bottom.pack(fill="x", pady=(8, 0))
        ttk.Button(live_bottom, text="Remover fonte", command=lambda: drop()).pack(side="left")
        ttk.Button(live_bottom, text="Verificar agora", command=sync_now).pack(side="left",
                                                                              padx=(8, 0))
        ttk.Label(live_bottom, textvariable=note, style="Hint.TLabel").pack(side="left",
                                                                            padx=(12, 0))

        base_table = ttk.Treeview(stored, show="headings", selectmode="extended", height=12,
                                  columns=("linha", "mes", "linhas", "planilha", "origem", "quando"))
        for column, caption, width in (("linha", "Linha", 70), ("mes", "Mês", 80),
                                       ("linhas", "Linhas", 70), ("planilha", "Planilha", 280),
                                       ("origem", "Origem", 220), ("quando", "Adicionado", 140)):
            base_table.heading(column, text=caption, anchor="w")
            base_table.column(column, width=width, anchor="w",
                              stretch=column in ("planilha", "origem"))
        base_table.pack(fill="both", expand=True)
        stored_bottom = ttk.Frame(stored, style="Card.TFrame")
        stored_bottom.pack(fill="x", pady=(8, 0))
        delete_button = ttk.Button(stored_bottom, text="Apagar selecionados",
                                   command=lambda: delete())
        delete_button.pack(side="left")
        ttk.Button(stored_bottom, text="Mostrar a base na tela", command=lambda: show_base()).pack(
            side="left", padx=(8, 0))
        base_note = tk.StringVar()
        ttk.Label(stored_bottom, textvariable=base_note, style="Hint.TLabel").pack(
            side="left", padx=(12, 0))

        def short(text):
            """A key fits the column: the end of a link or a path says which file it is."""
            text = str(text)
            return text if len(text) <= 42 else "…" + text[-41:]

        def fill():
            if not dialog.winfo_exists():
                return
            sources_table.delete(*sources_table.get_children())
            for source in load_sources()["sources"]:
                changed = str(source.get("changed") or "")
                ignored = list((source.get("ignored") or {}).values())
                if source.get("problem"):
                    situation = source["problem"]
                elif source_versions(source):
                    situation = "em dia · {} arquivo(s) · conferido {}{}".format(
                        source.get("files", 1), str(source.get("checked", ""))[11:16],
                        f" · ignorado: {ignored[0]}" if ignored else "")
                else:
                    situation = "aguardando a primeira leitura"
                sources_table.insert("", "end", iid=source["id"], values=(
                    source.get("label", ""), short(source["key"]),
                    f"{changed[8:10]}/{changed[5:7]} {changed[11:16]}" if changed else "—",
                    source.get("by") or "—", situation))
            note.set("confere a cada minuto enquanto o app está aberto"
                     + (" · verificando..." if book.get("syncing") else ""))
            base_table.delete(*base_table.get_children())
            for entry in sorted(base_entries(), key=lambda e: (str(e.get("line") or ""),
                                                               e.get("month") or [0, 0])):
                month = entry.get("month")
                base_table.insert("", "end", iid=entry["id"], values=(
                    entry.get("line") or "—", f"{month[1]:02d}/{month[0]}" if month else "?",
                    thousands(entry.get("rows", 0)), entry.get("file", ""),
                    entry.get("origin", ""), str(entry.get("added", "")).replace("T", " ")[:16]))
            base_note.set(f"{len(base_table.get_children())} planilhas na base")

        book["api_refresh"] = fill

        def save_token():
            set_token(token.get())
            say("Token salvo: as fontes são conferidas agora e a cada minuto.", dark["ok"])
            sync_now()
            fill()

        def probe():
            """Read only: what the typed token and key reach. Nothing is saved or downloaded."""
            typed_key, typed_token = key.get(), token.get()
            probe_button.state(["disabled"])
            note.set("Testando a conexão...")
            news = queue.Queue()

            def work():
                try:
                    news.put(("ok", probe_source(typed_key, typed_token)))
                except Exception as error:
                    news.put(("erro", error))

            threading.Thread(target=work, daemon=True).start()

            def poll():
                try:
                    kind, payload = news.get_nowait()
                except queue.Empty:
                    dialog.after(100, poll)
                    return
                probe_button.state(["!disabled"])
                note.set("")
                if kind == "erro":
                    messagebox.showerror("Testar conexão", f"Não conectou.\n\n{payload}",
                                         parent=dialog)
                    return
                shown = payload[:16] + ([f"... e mais {len(payload) - 16}"] if len(payload) > 16
                                        else [])
                messagebox.showinfo("Testar conexão", "\n".join(shown) + "\n\nNada foi baixado"
                                    " nem salvo. Para trazer os dados, use Adicionar fonte.",
                                    parent=dialog)

            dialog.after(100, poll)

        def add():
            try:
                add_source(label.get(), key.get())
            except ValueError as error:
                messagebox.showerror("Adicionar API", str(error), parent=dialog)
                return
            key.set("")
            sync_now()  # the first check is the first download
            fill()

        def drop():
            picked = sources_table.selection()
            if not picked:
                note.set("Selecione a fonte que vai sair.")
                return
            name = sources_table.set(picked[0], "mes")
            if messagebox.askyesno("Remover fonte", f"Parar de acompanhar {name}? O que ela já"
                                   " trouxe continua na base.", parent=dialog):
                remove_source(picked[0])
                fill()

        def feed_file():
            chosen = filedialog.askopenfilename(
                parent=dialog, initialdir=start_dir(),
                filetypes=[("Pacote ou Excel", "*.zip *.xlsx"), ("Pacote de relatórios", "*.zip")])
            if not chosen:
                return
            send_button.state(["disabled"])
            note.set("Lendo...")
            news = queue.Queue()

            def work():
                try:
                    data = load(chosen, progress=lambda i, n, name: news.put(
                        ("step", f"Lendo {i}/{n}: {name}")))
                    news.put(("ok", base_add(data, Path(chosen).name)))
                except Exception as error:
                    news.put(("erro", error))

            threading.Thread(target=work, daemon=True).start()

            def poll():
                try:
                    kind, payload = news.get_nowait()
                except queue.Empty:
                    dialog.after(100, poll)
                    return
                if kind == "step":
                    note.set(payload)
                    dialog.after(50, poll)
                    return
                send_button.state(["!disabled"])
                if kind == "erro":
                    note.set("")
                    messagebox.showerror("Adicionar API", str(payload), parent=dialog)
                    return
                saved, replaced = payload
                fill()
                show_base("{} planilha(s) salvas na base{}.".format(
                    len(saved), f", {len(replaced)} substituída(s)" if replaced else ""))

            dialog.after(100, poll)

        def delete():
            picked = base_table.selection()
            if not picked:
                base_note.set("Selecione o que apagar (Ctrl ou Shift marcam vários).")
                return
            names = [base_table.set(i, "planilha") for i in picked]
            shown = "\n".join(names[:12]) + ("\n..." if len(names) > 12 else "")
            if not messagebox.askyesno("Apagar da base", "Apagar {} planilha(s) da base para"
                                       " sempre?\n\n{}".format(len(picked), shown), parent=dialog):
                return
            base_remove(picked)
            fill()
            show_base(f"{len(picked)} planilha(s) apagadas da base.")

        fill()
        centre(dialog, (1000, 640))

    def as_date(text):
        if not text.strip():
            return None
        day = read_typed_date(text, book.get("month"))
        if day is None and book.get("month") and re.fullmatch(r"\d{1,2}", text.strip()):
            raise ValueError("O dia {} não existe em {:02d}/{}.".format(
                text.strip(), book["month"][1], book["month"][0]))
        if day is None:
            raise ValueError(f"Data {text.strip()!r} não dá para ler — use dd/mm/aaaa"
                             " ou só o dia.")
        return day

    def typing_date(field):
        """Put the slashes in while the user types, keeping the cursor at the end."""
        def on_key(event):
            if event.keysym in ("BackSpace", "Delete", "Left", "Right", "Tab", "Return"):
                return
            masked = mask_date(event.widget.get())
            if masked != event.widget.get():
                event.widget.delete(0, "end")
                event.widget.insert(0, masked)
            event.widget.icursor("end")
        field.bind("<KeyRelease>", on_key)

    def terms():
        return [t.strip() for t in keyword.get().split(",") if t.strip()]

    def redraw():
        draw_chips(chips, book["chosen"], drop_line)
        line_box["values"] = [n for n in book["lines"] if n not in book["chosen"]]
        line_choice.set("")

    def add_line(event=None):
        if event:
            event.widget.selection_clear()
        name = line_choice.get().strip()
        if name and name not in book["chosen"]:
            book["chosen"].append(name)
            redraw()
            refresh()

    def drop_line(name):
        book["chosen"].remove(name)
        redraw()
        refresh()

    def chosen_lines():
        return list(book["chosen"])

    def filters():
        return dict(text_col=book["text"], terms=terms(), date_col=book["date"],
                    start=as_date(start_date.get()), end=as_date(end_date.get()),
                    hour_col=book["hour"],
                    shift=next((n for n, label in SHIFT_LABELS.items() if label == hour.get()), None),
                    line_col=book.get("line", ""), lines=chosen_lines())

    def refresh(event=None):
        if event is not None and isinstance(getattr(event, "widget", None), ttk.Combobox):
            event.widget.selection_clear()
        if not book["rows"]:
            return
        try:
            chosen = filters()
            result = summarize(book["headers"], book["rows"], only_with_text=True, **chosen)
            # production exists in every hour, even the ones with nothing written down
            hours_result = summarize(book["headers"], book["rows"], **chosen)
        except Exception as error:
            say(f"Erro: {error}", dark["error"])
            return
        columns = result["columns"]
        when = [columns.index(c) for c in (book["date"], book["label"]) if c in columns]
        body = columns.index(book["text"]) if book["text"] in columns else None
        marks = [t.casefold() for t in terms()]

        line_at = columns.index(book["line"]) if book.get("line") in columns else None
        listing.configure(state="normal")
        listing.delete("1.0", "end")
        current = None
        for row in result["rows"]:
            # one block per production line, in the order the lines appear
            if line_at is not None and row[line_at] != current:
                current = row[line_at]
                listing.insert("end", f"\nLINHA {show(current)}\n", "line")
            listing.insert("end", " · ".join(show(row[i]) for i in when) + "\n", "when")
            text = show(row[body]).strip() if body is not None else ""
            listing.insert("end", sentences_with(text, marks) + "\n\n")
        if not result["rows"]:
            listing.insert("end", "Nenhuma ocorrência com esses filtros.\n", "none")
        for term in terms():
            at = "1.0"
            while at := listing.search(term, at, nocase=True, stopindex="end"):
                listing.tag_add("hit", at, at := f"{at}+{len(term)}c")
        listing.configure(state="disabled")

        draw_board(hours_result)

        summary.set("{} a {}".format(*map(show, result["period"])) if result["period"]
                    else "Nenhuma data no filtro")
        say(book.pop("just_loaded", None)
            or "Com palavra-chave, só a frase que a contém aparece. Enter refiltra.")

    def draw_board(result):
        """Produção tab: the same filter, seen as production numbers."""
        numbers = production_summary(book["headers"], result["matched"])
        book["board"] = numbers
        book["columns"], book["monthly"] = chart_columns(numbers)
        book["picked"] = None
        about = ", ".join(terms())
        chart_title.set("Tempo de troca (changeover) e produção por {}   ·   troca total {:.0f} min{}"
                        .format("mês" if book["monthly"] else "dia",
                                sum(v for _, v in numbers["trocas_por_dia"]),
                                f"   ·   só horas com {about}" if about else ""))
        sku_title.set(f"SKU e changeover por dia — {about}" if about
                      else "SKU e changeover por dia")
        tile_values["plan"].configure(text=thousands(numbers["plan"]))
        tile_values["produzido"].configure(text=thousands(numbers["produzido"]))
        tile_values["faltou"].configure(text=thousands(numbers["faltou"]),
                                        foreground=dark["miss"] if numbers["faltou"] > 0
                                        else dark["text"])
        tile_values["atingido"].configure(
            text="—" if numbers["atingido"] is None else f"{numbers['atingido'] * 100:.1f}%")
        tile_values["ocorrencias"].configure(
            text=thousands(sum(n for _, n in numbers["equipamentos"])))
        redraw_board()
        sku_list.configure(state="normal")
        sku_list.delete("1.0", "end")
        for day in numbers["skus_por_dia"]:
            skus, swaps = day["skus"], day["swaps"]
            # a mark per day, so a click on the chart can bring that day into view
            sku_list.mark_set(f"day{day['day']:%Y%m%d}", "end-1c")
            sku_list.mark_gravity(f"day{day['day']:%Y%m%d}", "left")
            head = show(day["day"]) + (f"   ·   {day['file']}" if day["file"] else "")
            sku_list.insert("end", head + "\n", "day")
            if not skus:
                sku_list.insert("end", "   sem SKU anotado\n", "quiet")
            else:
                sku_list.insert("end", "   ", "quiet")
                sku_list.insert("end", skus[0][0], "code")
                sku_list.insert("end", "   " + skus[0][1].split(" - ", 1)[-1] + "\n")
            # each changeover moves the day on to the next SKU written in the cell
            for i, (hours, said, took) in enumerate(swaps):
                when = ", ".join(hours) if hours else "hora não anotada"
                sku_list.insert("end", "   ⇄ houve um changeover {}{}\n".format(
                    when, f"  ({took})" if took else ""), "swap")
                if i + 1 < len(skus):
                    sku_list.insert("end", "      passa a ", "quiet")
                    sku_list.insert("end", skus[i + 1][0], "code")
                    sku_list.insert("end", "   " + skus[i + 1][1].split(" - ", 1)[-1] + "\n")
                elif skus:
                    sku_list.insert("end", "      a planilha não anotou o SKU seguinte\n", "quiet")
                sku_list.insert("end", f"      {said}\n")
            sku_list.insert("end", "\n")
        if not numbers["skus_por_dia"]:
            sku_list.insert("end", "Nenhum SKU no filtro.", "quiet")
        sku_list.configure(state="disabled")

    def redraw_board(*_):
        if book.get("board"):
            draw_combo_chart(day_chart, book["columns"], book["monthly"], book.get("picked"))
            draw_rank_chart(rank_chart, book["board"]["equipamentos"], highlight=book.get("over"),
                            colour=dark["miss"], strong="#e88a4a", ink=dark["text"],
                            quiet=dark["muted"])

    def slot_under(event):
        slot = tag_index(day_chart, event.x, event.y, "slot")
        return slot if slot is not None and slot < len(book.get("columns") or []) else None

    def pick_day(event):
        """Click a day: highlight it and bring its SKU and changeovers into view."""
        slot = slot_under(event)
        if slot is None:
            return
        book["picked"] = slot
        redraw_board()
        when = book["columns"][slot][0]
        days = [d["day"] for d in book["board"]["skus_por_dia"]
                if (d["day"].replace(day=1) if book["monthly"] else d["day"]) == when]
        sku_list.tag_remove("picked", "1.0", "end")
        if days:
            mark = f"day{days[0]:%Y%m%d}"
            sku_list.tag_add("picked", mark, f"{mark} lineend+1c")
            sku_list.yview(mark)
        say("{} selecionado. Duplo clique filtra só ele.".format(
            when.strftime("%m/%Y") if book["monthly"] else show(when)))

    def zoom_in(event):
        """Double-click a day (or month) to filter the whole screen down to it."""
        slot = slot_under(event)
        if slot is None:
            return
        day = book["columns"][slot][0]
        if book["monthly"]:
            first = day.replace(day=1)
            last = first.replace(day=calendar.monthrange(first.year, first.month)[1])
        else:
            first = last = day
        start_date.set(show(first))
        end_date.set(show(last))
        refresh()
        say("Abrindo {}. Apague as datas para voltar ao período inteiro.".format(
            first.strftime("%m/%Y") if book["monthly"] else show(first)))

    def day_hover(event):
        slot = slot_under(event)
        if slot is None:
            day_chart.delete("tip")
            return
        when, minutes, planned, produced = book["columns"][slot]
        lines = [when.strftime("%m/%Y") if book["monthly"] else show(when),
                 f"Changeover: {minutes:.0f} min", f"Produzido: {thousands(produced)}"]
        if planned is None:
            lines.append("Plano: não anotado")
        else:
            gap = produced - planned
            lines += [f"Plano: {thousands(planned)}",
                      "{}{} ({:.0f}%)".format("+" if gap >= 0 else "", thousands(gap),
                                              produced / planned * 100 if planned else 0)]
        draw_tip(day_chart, event.x, event.y, lines)

    def machine_under(event):
        """Which machine the mouse is on."""
        return tag_index(rank_chart, event.x, event.y, "bar")

    def rank_hover(event):
        over = machine_under(event)
        if over == book.get("over"):
            return
        book["over"] = over
        rank_chart.configure(cursor="hand2" if over is not None else "")
        redraw_board()
        if over is not None:
            name = book["board"]["equipamentos"][over][0]
            say(f"Clique em {name} para ver só o que envolve esse equipamento.")

    def rank_click(event):
        """Clicking a machine takes the user to its occurrences, already filtered."""
        over = machine_under(event)
        if over is None or not book.get("board"):
            return
        name = book["board"]["equipamentos"][over][0]
        keyword.set(name)
        refresh()
        tabs.select(frame)
        listing.see("1.0")
        say(f"Ocorrências de {name}. Apague a palavra-chave para ver todas de novo.")

    def download_file():
        """Linha, Data, Hora and Ocorrência of what the Ocorrências tab is showing."""
        if not book["rows"]:
            return
        try:
            chosen = filters()
            found = summarize(book["headers"], book["rows"], only_with_text=True, **chosen)
            target = filedialog.asksaveasfilename(
                defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                initialfile="ocorrencias.xlsx", initialdir=start_dir())
            if not target:
                return
            written = download(target, book["headers"], found["matched"], book.get("line", ""),
                               book["date"], book["label"] or book["hour"], book["text"],
                               chosen["terms"])
            say(f"{written} ocorrências gravadas em {target}", dark["ok"])
        except Exception as error:
            say(f"Erro: {error}", dark["error"])

    show_history()
    day_chart.bind("<Configure>", redraw_board)
    day_chart.bind("<Motion>", day_hover)
    day_chart.bind("<Button-1>", pick_day)
    day_chart.bind("<Double-Button-1>", zoom_in)
    day_chart.bind("<Leave>", lambda _: day_chart.delete("tip"))
    rank_chart.bind("<Configure>", redraw_board)
    rank_chart.bind("<Motion>", rank_hover)
    rank_chart.bind("<Leave>", lambda _: (book.update(over=None), redraw_board()))
    rank_chart.bind("<Button-1>", rank_click)
    typing_date(first)
    typing_date(last)
    root.bind("<Return>", refresh)
    hour_box.bind("<<ComboboxSelected>>", refresh)
    line_box.bind("<<ComboboxSelected>>", add_line)
    apply_role()  # everyone starts as a viewer; the Unilever button logs the admin in
    if path:
        root.after(50, lambda: open_file(path))
    else:
        root.after(50, show_base)
    root.after(2000, keep_syncing)  # live sources: checked now and every minute after
    root.mainloop()


if __name__ == "__main__":
    if sys.argv[1:2] == ["--admin"] and len(sys.argv) == 3:
        # python excel_filter.py --admin <login>: (re)writes admin.json, asking the password twice
        import getpass

        typed = getpass.getpass("Senha do admin: ")
        if typed != getpass.getpass("De novo: "):
            sys.exit("As senhas não conferem.")
        print(f"Gravado em {make_admin(sys.argv[2], typed)}")
    else:
        main(*sys.argv[1:2])  # optional: path of a spreadsheet to open right away
